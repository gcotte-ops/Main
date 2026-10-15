"""Tests hors ligne de la détection des doublons (entreprises et contacts) et des filets de
sécurité ajoutés aux scripts de fusion. Les fiches de référence reprennent les groupes de
l'ancien classeur doublons_potentiels_entreprises_hubspot.xlsx (G-001, G-084, G-187...)."""

import contextlib
import csv
import datetime
import io
import itertools
import os
import random
import tempfile
import unittest

import detection_contacts_hubspot as dc
import detection_entreprises_hubspot as de
import familles
import fusion_doublons_hubspot as fc
import fusion_entreprises_hubspot as fe
import normalisation as nz

CERTAIN, PROBABLE = de.CERTAIN, de.PROBABLE


def ent(id, nom, ville="", siren="", siret="", tel="", domaine="", decideur="", contacts=0,
        creation="2021-07-09 14:48", transactions=None):
    return de.FicheEntreprise(id, nom=nom, ville=ville, siren=siren, siret=siret, telephone=tel,
                              domaine=domaine, type_de_decideur=decideur, nb_contacts=contacts,
                              creation=creation, nb_transactions=transactions)


# Groupes « Certain » de l'ancien classeur (valeurs telles qu'exportées).
REFERENCES = {
    "G-001": [ent("14939896642", "ABBAYE DE LA COUDRE", "Laval Cedex", "352761043",
                  "35276104300012", "332 43 02 85 85", "abbaye-coudre.fr"),
              ent("6543335692", "EURL ABBAYE DE LA COUDRE", "LAVAL CEDEX",
                  tel="00 00 00 00 00", domaine="4313.co")],
    "G-003": [ent("6542856282", "ABBAYE NOTRE DAME DES GARDES", "SAINT GEORGES DES GARDES",
                  tel="+332 41 29 57 10", domaine="www.abbayedesgardes.fr", contacts=3),
              ent("14938411806", "abbayedesgardes.fr", tel="332 41 29 57 10",
                  domaine="abbayedesgardes.fr", creation="2023-03-03 17:33"),
              ent("6543341847", "SARL DE L'ABBAYE NOTRE DAME DES GARDES",
                  "SAINT GEORGES DES GARDES", tel="00 00 00 00 00", domaine="3987.co")],
    "G-019": [ent("421424505038", "APM 22", "Saint-Brieuc", tel="+33296681680",
                  domaine="apm22.fr", contacts=2, creation="2026-03-13 10:53"),
              ent("421272184057", "APM22", "SAINT-BRIEUC", "379740913", "37974091300031",
                  contacts=2, creation="2026-03-13 10:53")],
    "G-031": [ent("421254450409", "Association Filéas", "Vitre", tel="+33299753123",
                  domaine="association-fileas.fr", contacts=1),
              ent("421335282925", "Association FILEAS", "VITRE", "305349466", "30534946600022",
                  contacts=1),
              ent("421290298555", "Association FILEAS", "RETIERS", "378206452",
                  "37820645200013", contacts=1)],
    "G-037": [ent("9542641081", "Sainte Marie Antony", "Antony", "301546503", "30154650300014",
                  "33146115800", "www.saintemarieantony.fr", "Ne sait pas encore", 2),
              ent("9542385249", "Association Sainte-Marie d'Antony", "Antony", "529572703",
                  "52957270300013", "0146115800", "saintemarieantony.fr", contacts=2),
              ent("6543335990", "ASSOCIATION SAINTE MARIE", "ANTONY", tel="00 00 00 00 00",
                  domaine="9185.co")],
    "G-040": [ent("281120678136", "Cabinet ballu", siren="351864384", siret="35186438400034",
                  tel="+33147753031", domaine="cabinet-ballu.fr", contacts=1),
              ent("282918297842", "Ballu", tel="+33147753031", domaine="ballu.fr", contacts=1),
              ent("32411987303", "cabinet ballu", "Puteaux", "449830660", contacts=1)],
    "G-082": [ent("15302177367", "COVEDIA IMMOBILIER", "HYERES", "830011532", "83001153200017",
                  "33 1 55 50 60 00", "https://www.covea.com/fr/covea-immobilier", contacts=3),
              ent("189174986987", "COVEA Immobilier", "HYERES", "797908654", "79790865400024",
                  "+33 1 55 50 60 00", "covea.com", contacts=21),
              ent("183282788555", "Covéa Immobilier", domaine="covea-immobilier.fr",
                  contacts=5)],
    "G-084": [ent("440764019935", "DEPARTEMENT DE LA LOIRE-ATLANTIQUE", "NANTES", "224400028",
                  "22440002800011", "+33240991000", "loire-atlantique.fr", contacts=1),
              ent("18058650763", "Loire-atlantique.fr", "Nantes", tel="+33 2 40 71 03 50",
                  domaine="loire-atlantique.fr", contacts=1),
              ent("445112249553", "Les Services de l'État en Loire-Atlantique",
                  domaine="loire-atlantique.gouv.fr", contacts=1)],
    "G-103": [ent("15746638500", "EHPAD CH ALBERT", "ALBERT CEDEX", "268000031",
                  "268 000 031 00032", "+33322744147", "ch-albert.fr", "Ne sait pas encore", 1),
              ent("15747049734", "EHPAD PCP ALBERT", "ALBERT", siret="395 135 098 00014",
                  tel="+33322744147")],
    "G-160": [ent("6543220604", "INSTITUTION SAINT LOUIS DE GONZAGUE", "PERPIGNAN", "776188005",
                  "77618800500016", "04 68 64 17 71", "https://gonzague.com/",
                  "Ne sait pas encore", 5),
              ent("9638282589", "Saint Louis de Gonzague", "Perpignan", "812489318",
                  "81248931800012", "+33 4 68 61 23 91", "gonzague.com", contacts=1),
              ent("447779243208", "", domaine="slnd.gonzague.com", contacts=1)],
    "G-187": [ent("6543371629", "OGEC MARMOUTIER", "TOURS CEDEX", "308050418",
                  tel="33247883535", domaine="2770.co", contacts=1),
              ent("7436392842", "marmoutier.org", "Tours", tel="+33247883535",
                  domaine="marmoutier.org", contacts=2),
              ent("188501102817", "Mairie - Marmoutier", siren="216702837",
                  siret="21670283700057", domaine="marmoutier.fr", contacts=1)],
    "G-227": [ent("16328405384", "Ville de Saint Mathieu de Tréviers",
                  "Saint-Mathieu-de-Tréviers", tel="+33467552028",
                  domaine="ville-saint-mathieu-de-treviers.fr"),
              ent("30477778306", "Ville de Saint Mathieu de Tréviers",
                  "Saint-Mathieu-de-Tréviers", domaine="villesmdt.fr", contacts=1)],
}


def analyser(fiches=None, **options):
    fiches = fiches if fiches is not None else [f for v in REFERENCES.values() for f in v]
    return de.Analyse(fiches, journal=lambda *a: None, **options).executer()


def ids(analyse, groupe):
    return {analyse.fiches[m].id for m in groupe.membres}


def groupes_de(analyse, *id_fiches, niveau=None):
    return [g for g in analyse.doublons if set(id_fiches) <= ids(analyse, g)
            and (niveau is None or g.niveau == niveau)]


def paires_distinctes(analyse):
    return {frozenset(ids(analyse, g)) for g in analyse.relecture["bloquees"]}


# --------------------------------------------------------------------------- #
# Normalisation (section 4)
# --------------------------------------------------------------------------- #

class Normalisation(unittest.TestCase):

    def test_valeurs_neutralisees(self):
        for domaine in ("4313.co", "https://www.4313.co/", "xn--manrse-6ua.com",
                        "sacrecoeur.paysdelaloire.e-lyco.fr", "monsite.wixsite.com/ecole",
                        "sites.google.com/view/x", "facebook.com/abbaye"):
            self.assertEqual(nz.hote_utile(domaine), "", domaine)
        self.assertEqual(nz.hote_utile("carrefree.fr"), "carrefree.fr")   # pas « free.fr »
        self.assertEqual(nz.telephone("00 00 00 00 00"), "")
        self.assertEqual(nz.nettoyer('="012345678"'), "012345678")
        for vide in ("Ne sait pas encore", "NC", "inconnu", "n/a", "-"):
            self.assertEqual(nz.nettoyer(vide), "")

    def test_noms(self):
        nom = "SARL de l'Abbaye Notre-Dame des Gardes (ancien site)"
        self.assertEqual(nz.nom_norm(nom), "sarl de l abbaye notre dame des gardes ancien site")
        self.assertEqual(nz.nom_sans_parentheses(nom), "sarl de l abbaye notre dame des gardes")
        self.assertEqual(nz.nom_equivalent(nom), "abbaye notre dame gardes")
        self.assertEqual(nz.nom_equivalent("OGEC École Sainte-Anne"), "sainte anne")
        self.assertEqual(nz.nom_equivalent("Groupe Scolaire St Jean"), "saint jean")
        self.assertEqual(nz.nom_norm("Congrégation Sœurs"), "congregation soeurs")
        self.assertEqual(nz.nom_distinctif(nz.nom_equivalent("Mairie de Nantes"),
                                           nz.ville("NANTES CEDEX")), "mairie")
        self.assertTrue(nz.nom_est_domaine("abbayedesgardes.fr"))
        self.assertTrue(nz.nom_est_domaine("www.marmoutier.org"))
        self.assertFalse(nz.nom_est_domaine("Abbaye des Gardes"))

    def test_nom_generique(self):
        a = ent("1", "Mairie de Nantes", "Nantes")
        b = ent("2", "Abbaye de la Coudre", "Laval")
        self.assertTrue(a.generique)               # nom distinctif « mairie » : 6 lettres
        self.assertFalse(b.generique)
        # Même nom équivalent dans 3 villes différentes : générique dans toute la base.
        fiches = [ent(str(i), "École Saint Joseph", v) for i, v in
                  enumerate(("Nantes", "Angers", "Laval"))] + [ent("9", "Saint Joseph", "Laval")]
        analyse = de.Analyse(fiches, journal=lambda *a: None)
        self.assertTrue(all(f.generique for f in analyse.fiches))

    def test_villes(self):
        self.assertEqual(nz.ville("Lyon 6Eme"), "lyon")
        self.assertEqual(nz.ville("Lyon-9E-Arrondissement"), "lyon")
        self.assertEqual(nz.ville("LAVAL CEDEX 9"), "laval")
        self.assertEqual(nz.ville("ST JUST EN CHAUSSEE"), "saint just en chaussee")
        self.assertEqual(nz.ville("SAINT ANDRE LEZ LILLE"), nz.ville("Saint-André-lès-Lille"))
        self.assertTrue(nz.villes_identiques("segre", nz.ville("Segré-en-Anjou Bleu")))
        self.assertFalse(nz.villes_identiques("segre", "segrez"))

    def test_telephones(self):
        self.assertEqual(nz.telephone("+33 (0)2 41 29 57 10"), "0241295710")
        self.assertEqual(nz.telephone("332 41 29 57 10"), "0241295710")
        self.assertEqual(nz.telephone("+1 650-618-7714"), "")

    def test_domaines(self):
        self.assertEqual(nz.hote("https://www.etu.univ-amu.fr/x"), "etu.univ-amu.fr")
        self.assertEqual(nz.domaine_principal("etu.univ-amu.fr"), "univ-amu.fr")
        self.assertEqual(nz.domaine_principal("loire-atlantique.gouv.fr"),
                         "loire-atlantique.gouv.fr")
        self.assertEqual(nz.domaine_principal("lycee-x.ac-nantes.fr"), "lycee-x.ac-nantes.fr")
        self.assertEqual(nz.domaine_principal("ac-nantes.fr"), "")      # suffixe seul
        self.assertEqual(nz.domaine_principal("site.asso.fr"), "site.asso.fr")
        self.assertEqual(nz.comparer_domaines("abbaye.fr", "abbaye.fr"),
                         {"meme_domaine", "meme_hote"})
        self.assertEqual(nz.comparer_domaines("loire-atlantique.fr", "loire-atlantique.gouv.fr"),
                         {"domaines_differents", "domaine_proche"})
        # Sous-domaines différents : domaine proche seulement, jamais un lien fort.
        self.assertEqual(nz.comparer_domaines("gonzague.com", "slnd.gonzague.com"),
                         {"domaine_proche"})
        self.assertEqual(nz.comparer_domaines("celc.caisse-epargne.fr", "cecaz.caisse-epargne.fr"),
                         {"domaine_proche"})
        self.assertEqual(nz.comparer_domaines("ballu.fr", "cabinet-ballu.fr"),
                         {"domaines_differents", "domaine_proche"})

    def test_siren(self):
        self.assertEqual(nz.siren_depuis("352 761 043", ""), "352761043")
        self.assertEqual(nz.siren_depuis("", "395 135 098 00014"), "395135098")
        self.assertEqual(nz.siren_depuis("123", "456"), "")
        f = fe.Entreprise(1, {"siret": "39513509800014"})
        self.assertEqual(fe.siren_de(f), "395135098")

    def test_emails(self):
        self.assertTrue(nz.est_webmail("x@gmail.com"))
        self.assertFalse(nz.est_webmail("x@alter-watt.com"))
        for adresse in ("contact@x.fr", "info.ecole@x.fr", "accueil-2@x.fr", "no-reply@x.fr",
                        "secretariat@x.fr"):
            self.assertTrue(nz.est_generique(adresse), adresse)
        for adresse in ("contacts@x.fr", "jean.dupont@x.fr", "informatique@x.fr"):
            self.assertFalse(nz.est_generique(adresse), adresse)
        self.assertEqual(nz.emails_liste("A@x.fr; b@y.fr"), ["a@x.fr", "b@y.fr"])


# --------------------------------------------------------------------------- #
# Règle des familles (section 7, test d'acceptation 1)
# --------------------------------------------------------------------------- #

class Familles(unittest.TestCase):

    CAS = [("OGEC Marmoutier", "Mairie - Marmoutier", "activité"),
           ("Département de la Loire-Atlantique", "Les Services de l'État en Loire-Atlantique",
            "activité"),
           ("Foncière des Régions", "Association Les Régions", "nature"),
           ("Association diocésaine de Reims",
            "Association immobilière catholique (immobilier scolaire)", "activité"),
           ("Fondation Alix Le Clerc", "Association Alix Le Clerc", ""),
           ("OGEC Sainte Marie", "École Sainte Marie", ""),
           ("OGEC Sainte Marie", "Association Sainte Marie", ""),
           ("EURL Abbaye de la Coudre", "Abbaye de la Coudre", ""),
           ("Maison diocésaine de Tarentaise", "Association diocésaine de Tarentaise", ""),
           ("Institut Pasteur", "Institut Bossuet", "")]

    def test_cas_de_reference(self):
        for a, b, attendu in self.CAS:
            motif = familles.incompatibilite(a, b)
            if attendu:
                self.assertIn(attendu, motif, (a, b))
            else:
                self.assertEqual(motif, "", (a, b))
            self.assertEqual(bool(motif), bool(familles.incompatibilite(b, a)))

    def test_details(self):
        # Nature explicite plutôt qu'implicite.
        self.assertEqual(familles.etiquettes("Association immobilière catholique"),
                         ({"immobilier"}, {"asso_fondation"}))
        # Position « debut » : « région » seulement en tête de nom.
        self.assertEqual(familles.etiquettes("Foncière des Régions")[0], {"immobilier"})
        self.assertEqual(familles.etiquettes("Région Pays de la Loire")[0], {"collectivite"})
        # Pluriel toléré, mots ambigus jamais seuls.
        self.assertEqual(familles.etiquettes("Écoles Saint-Jean")[0], {"enseignement"})
        self.assertEqual(familles.etiquettes("Résidence du Parc"), (set(), set()))
        self.assertEqual(familles.etiquettes("Résidence autonomie du Parc")[0], {"medico_social"})

    def test_lexique_editable(self):
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "lexique.csv")
        with open(chemin, "w", encoding="utf-8") as f:
            f.write("dimension;famille;marqueur;position;nature_implicite\n"
                    "activite;sport;club;;asso_fondation\n"
                    "nature;societe;sarl;;\n"
                    "activite;sante;maison;;\n")                 # ambigu : ignoré
        lexique = familles.charger_lexique(chemin, avertir=False)
        self.assertEqual([m.marqueur for m in lexique], ["club", "sarl"])
        self.assertTrue(familles.incompatibilite("Club de Nantes", "SARL de Nantes", lexique))


# --------------------------------------------------------------------------- #
# Paires d'entreprises (section 5)
# --------------------------------------------------------------------------- #

class PairesEntreprises(unittest.TestCase):

    def niveau(self, a, b):
        return de.evaluer(a, b)

    def test_c1_et_contradictions(self):
        a = ent("1", "Abbaye de la Coudre", "Laval")
        b = ent("2", "EURL ABBAYE DE LA COUDRE", "LAVAL CEDEX", tel="00 00 00 00 00")
        p = self.niveau(a, b)
        self.assertEqual((p.niveau, p.regles), (CERTAIN, ["C1"]))
        self.assertEqual(de.raison(p), "Nom équivalent (hors type d’établissement) + même ville")
        b.ville_norm = "angers"
        self.assertEqual(self.niveau(a, b).niveau, "")        # villes différentes, aucun lien fort

    def test_c6_domaines_differents(self):
        a, b = REFERENCES["G-227"]
        p = self.niveau(a, b)
        self.assertEqual(p.niveau, PROBABLE)                  # un seul lien fort
        self.assertIn("domaines différents", de.raison(p))
        b.tel = a.tel                                         # deux liens forts : C6
        self.assertEqual(self.niveau(a, b).niveau, CERTAIN)

    def test_c3_fiche_nommee_par_son_domaine(self):
        a, b, _ = REFERENCES["G-003"]
        self.assertEqual(self.niveau(a, b).regles, ["C3"])
        c = ent("3", "abbayedesgardes.fr", tel="0241295710", domaine="www.abbayedesgardes.fr/x")
        self.assertEqual(self.niveau(a, c).niveau, CERTAIN)   # même hôte
        d = ent("4", "", domaine="intranet.abbayedesgardes.fr", tel="0241295710")
        self.assertEqual(self.niveau(a, d).niveau, PROBABLE)  # autre hôte : domaine proche

    def test_c4_siren_centralise(self):
        a = ent("1", "Adef Résidences", "Ivry", "323649525", tel="01 46 70 16 02",
                decideur="Centralisé")
        b = ent("2", "Fondation ADEF Ivry", "Ivry", siret="32364952500686", tel="01 46 70 16 02",
                decideur="Centralisé")
        self.assertEqual(self.niveau(a, b).regles, ["C4"])
        self.assertEqual(self.niveau(a, b).niveau, CERTAIN)
        # Les exclusions de 5.4 valent pour toutes les règles : nom inclus -> Probable.
        c = ent("3", "EHPAD ADEF", "Ivry", siret="32364952500686", tel="01 46 70 16 02",
                decideur="Centralisé")
        self.assertEqual((self.niveau(a, c).regles, self.niveau(a, c).niveau), (["C4"], PROBABLE))
        b.centralise = False                                  # SIREN commun sans Centralisé
        p = self.niveau(a, b)
        self.assertEqual((p.niveau, p.categorie), ("", "siren"))

    def test_c5_faute_et_siren(self):
        a = ent("1", "Institution Notre Dame", "Laval", "123456789")
        b = ent("2", "Institution Notre Damme", "Laval", "123456789")
        self.assertEqual(self.niveau(a, b).niveau, CERTAIN)
        b.siren = ""
        self.assertEqual(self.niveau(a, b).niveau, PROBABLE)  # faute sans SIREN identique

    def test_jamais_certain(self):
        # Nom inclus, même avec deux liens forts.
        a = ent("1", "Le Grand Narbonne", "Narbonne", tel="0468581458", domaine="legrandnarbonne.com")
        b = ent("2", "Communauté d'agglomération - Le Grand Narbonne", "Narbonne",
                tel="+334 68 58 14 58", domaine="legrandnarbonne.com")
        self.assertEqual(self.niveau(a, b).niveau, PROBABLE)
        # Ville absente sans autre lien fort.
        a = ent("1", "Almaviva Santé", "Marseille", tel="0491210606")
        b = ent("2", "ALMAVIVA SANTE")
        self.assertEqual(self.niveau(a, b).niveau, PROBABLE)

    def test_nom_generique(self):
        a = ent("1", "Mairie de Gap", "Gap", tel="0492510000")
        b = ent("2", "MAIRIE DE GAP", "GAP")
        self.assertEqual(self.niveau(a, b).niveau, PROBABLE)  # C2 : deux liens forts exigés
        b.tel = a.tel
        self.assertEqual(self.niveau(a, b).regles, ["C2"])

    def test_probables(self):
        # Fiche nommée par son domaine + même domaine, sans ville ni téléphone communs.
        a = ent("1", "Abbaye des Gardes", "Saint Georges", domaine="abbayedesgardes.fr")
        b = ent("2", "abbayedesgardes.fr", domaine="abbayedesgardes.fr")
        self.assertEqual(self.niveau(a, b).niveau, PROBABLE)
        # Noms différents + même domaine + même téléphone : changement de nom ?
        a = ent("1", "Carcassonne Agglo", "Carcassonne", tel="0468105600",
                domaine="carcassonne-agglo.fr")
        b = ent("2", "Agglomération - Carcassonne", "Carcassonne", tel="0468105600",
                domaine="carcassonne-agglo.fr")
        p = self.niveau(a, b)
        self.assertEqual(p.niveau, PROBABLE)
        self.assertIn("changement de nom", de.raison(p))

    def test_blocages(self):
        a, b, _ = REFERENCES["G-160"]
        p = self.niveau(a, b)
        self.assertTrue(p.bloquee)
        self.assertEqual(p.niveau, "")
        self.assertIn("B1", p.blocage[0])
        a, _, c = REFERENCES["G-084"]
        self.assertIn("B2", self.niveau(a, c).blocage[0])

    def test_filiale_et_lieux(self):
        contexte = de.Contexte(lieux={"paris", "angers"})
        a = ent("1", "ESSCA Angers", "Angers", domaine="essca.fr")
        b = ent("2", "ESSCA", "Paris", domaine="essca.fr")
        p = de.evaluer(a, b, contexte)
        self.assertEqual((p.niveau, p.categorie), ("", "filiale"))
        a = ent("1", "ESSCA - Paris", "Paris", domaine="essca.fr")
        b = ent("2", "ESSCA - Angers", "Angers", domaine="essca.fr")
        self.assertIn("lieux différents dans les noms", de.liens_texte(de.evaluer(a, b, contexte)))

    def test_homonymes(self):
        a, _, c = REFERENCES["G-031"]
        self.assertEqual(de.evaluer(a, c).categorie, "homonyme")


# --------------------------------------------------------------------------- #
# Tests d'acceptation (section 10) sur les groupes de l'ancien classeur
# --------------------------------------------------------------------------- #

class Acceptation(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.analyse = analyser()

    def test_2_departement_et_services_de_l_etat(self):
        dep, _, etat = (f.id for f in REFERENCES["G-084"])
        self.assertFalse(groupes_de(self.analyse, dep, etat, niveau=CERTAIN))
        self.assertFalse(groupes_de(self.analyse, dep, etat))
        self.assertIn(frozenset((dep, etat)), paires_distinctes(self.analyse))

    def test_3_ogec_et_mairie_marmoutier(self):
        ogec, _, mairie = (f.id for f in REFERENCES["G-187"])
        self.assertFalse(groupes_de(self.analyse, ogec, mairie))
        self.assertIn(frozenset((ogec, mairie)), paires_distinctes(self.analyse))

    def test_4_aucun_groupe_certain_avec_deux_siren(self):
        for g in self.analyse.doublons:
            sirens = {self.analyse.fiches[m].siren for m in g.membres} - {""}
            self.assertLessEqual(len(sirens), 1, g.numero)
        for ref in ("G-031", "G-037", "G-040", "G-082", "G-103", "G-160", "G-187"):
            sirens = {f.siren for f in REFERENCES[ref]} - {""}
            for a, b in itertools.combinations(sirens, 2):
                ids_a = [f.id for f in REFERENCES[ref] if f.siren == a]
                ids_b = [f.id for f in REFERENCES[ref] if f.siren == b]
                self.assertFalse(groupes_de(self.analyse, ids_a[0], ids_b[0], niveau=CERTAIN))

    def test_5_domaine_proche(self):
        institution, _, slnd = (f.id for f in REFERENCES["G-160"])
        self.assertFalse(groupes_de(self.analyse, institution, slnd, niveau=CERTAIN))
        self.assertTrue(groupes_de(self.analyse, institution, slnd, niveau=PROBABLE))

    def test_6_aucun_chainage_dans_un_groupe_certain(self):
        analyse = analyser(REFERENCES["G-003"] + [
            # Chaîne A - B - C : même téléphone A/B, même domaine B/C, A et C sans lien direct.
            ent("1", "Lycée Saint Louis", "Brest", tel="0298000000"),
            ent("2", "Lycee Saint-Louis", "Brest", tel="02 98 00 00 00", domaine="stlouis.fr"),
            ent("3", "lycée saint louis", "", domaine="stlouis.fr")])
        for g in analyse.doublons:
            if g.niveau == CERTAIN:
                for x, y in itertools.combinations(g.membres, 2):
                    self.assertEqual(analyse.paire(x, y).niveau, CERTAIN, g.numero)
        chaines = [g for g in analyse.doublons if g.origine is not None]
        self.assertTrue(chaines)
        for g in chaines:
            self.assertEqual(g.niveau, PROBABLE)
            self.assertIs(g.racine, g.origine.racine)

    def test_7_non_regression(self):
        a = self.analyse
        self.assertTrue(groupes_de(a, *(f.id for f in REFERENCES["G-001"]), niveau=CERTAIN))
        # G-003 : les deux fiches liées entre elles restent Certain ; la SARL n'a aucun lien
        # direct avec la fiche nommée par son domaine : groupe Probable « lien indirect ».
        abbaye, domaine, sarl = (f.id for f in REFERENCES["G-003"])
        certain = groupes_de(a, abbaye, domaine, niveau=CERTAIN)
        self.assertEqual(len(certain), 1)
        chaine = groupes_de(a, abbaye, sarl, niveau=PROBABLE)
        self.assertEqual(chaine[0].origine, certain[0])
        lignes = a.lignes()["Doublons"]
        motif = next(l["Raison du rapprochement"] for l in lignes if l["Groupe"] == chaine[0].numero)
        self.assertIn("Lien indirect (chaînage) avec " + certain[0].numero, motif)
        # G-227 : un seul lien fort et des domaines différents -> Probable (C6).
        self.assertTrue(groupes_de(a, *(f.id for f in REFERENCES["G-227"]), niveau=PROBABLE))
        # G-019 « APM 22 » : nom distinctif de 6 caractères ou moins, donc générique (4.2) ;
        # un seul lien fort (même ville) -> Probable selon C2.
        self.assertTrue(groupes_de(a, *(f.id for f in REFERENCES["G-019"]), niveau=PROBABLE))

    def test_racine_suggeree(self):
        a = self.analyse
        g = groupes_de(a, *[f.id for f in REFERENCES["G-003"]][:2], niveau=CERTAIN)[0]
        self.assertEqual(a.fiches[g.racine].id, "6542856282")    # le plus de contacts
        # Une transaction l'emporte sur les contacts ; la GT 3.0 départage les transactions.
        x = ent("1", "Abbaye X", "Laval", tel="0243000000", contacts=9, transactions=1)
        y = ent("2", "ABBAYE X", "LAVAL", tel="0243000000", contacts=1, transactions=2)
        z = ent("3", "Abbaye X", "Laval", tel="0243000000", contacts=0)
        analyse = analyser([x, y, z], gt3={"2"})
        self.assertEqual(analyse.fiches[analyse.doublons[0].racine].id, "2")
        analyse = analyser([x, y, z])
        self.assertEqual(analyse.fiches[analyse.doublons[0].racine].id, "1")
        # Alignée sur la fusion.
        fusion = [fe.Entreprise(f.id, {"name": f.nom}, nb_contacts=f.nb_contacts,
                                creation=f.creation, nb_transactions=f.nb_transactions)
                  for f in (x, y, z)]
        self.assertEqual(fe.choisir_racine(fusion, {"2"})[0].id, "2")
        self.assertEqual(fe.choisir_racine(fusion)[0].id, "1")

    def ecrire(self, analyse, chemin):
        lignes = analyse.lignes()
        onglets = [("Doublons", de.COLONNES, lignes["Doublons"])]
        onglets += [(o, de.COLONNES, lignes[o]) for _, _, o, _ in de.ONGLETS_RELECTURE]
        horodatage = datetime.datetime(2026, 10, 15, 12, 0)
        de.ecrire_classeur(chemin, de.synthese(analyse, "test", horodatage), onglets,
                           horodatage)
        return lignes

    def test_8_compatibilite_avec_la_fusion(self):
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "doublons_entreprises.xlsx")
        lignes = self.ecrire(self.analyse, chemin)
        ligne = lignes["Doublons"][0]
        self.assertEqual([k for k in ligne if not k.startswith("_")], de.COLONNES)
        self.assertEqual(de.COLONNES[:16], fe_entetes())
        lues = fe.lire_liste(chemin)
        self.assertEqual({l["ID HubSpot"] for l in lues},
                         {l["ID HubSpot"] for l in lignes["Doublons"] if l["Niveau"] == CERTAIN})
        rapport = os.path.join(dossier, "rapport.csv")
        compteur = fe.traiter(chemin, rapport, journal=lambda *a: 0)
        certains = sum(1 for g in self.analyse.doublons if g.niveau == CERTAIN)
        self.assertEqual(compteur, {"SIMULATION": certains})
        # La fusion retient la racine suggérée.
        racines = {l["Groupe"]: l["ID HubSpot"] for l in lignes["Doublons"]
                   if l["Racine suggérée"] == "Oui"}
        with open(rapport, encoding="utf-8-sig") as f:
            for l in csv.DictReader(f, delimiter=";"):
                self.assertEqual(l["Fiche racine"], racines[l["Groupe"]])

    def test_9_stabilite(self):
        fiches = [f for v in REFERENCES.values() for f in v]
        melangees = list(fiches)
        random.Random(4).shuffle(melangees)
        self.assertEqual(analyser(fiches).lignes(), analyser(melangees).lignes())
        dossier = tempfile.mkdtemp()
        chemins = [os.path.join(dossier, n) for n in ("a.xlsx", "b.xlsx")]
        self.ecrire(analyser(fiches), chemins[0])
        self.ecrire(analyser(melangees), chemins[1])
        contenus = [{f: fe.lire_xlsx(c, f) for f in de.feuilles_xlsx(c)} for c in chemins]
        self.assertEqual(contenus[0], contenus[1])

    def test_classeur(self):
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "doublons.xlsx")
        self.ecrire(self.analyse, chemin)
        self.assertEqual(de.feuilles_xlsx(chemin),
                         ["Synthèse", "Doublons", "Filiales et établissements",
                          "Homonymes à vérifier", "SIREN commun", "Entités distinctes"])
        lignes = fe.lire_xlsx(chemin)
        self.assertTrue(all(l["Lien HubSpot"] == "Ouvrir" for l in lignes))
        self.assertIn("14939896642", {l["ID HubSpot"] for l in lignes})     # ID en texte
        for groupe in {l["Groupe"] for l in lignes}:
            self.assertEqual(sum(1 for l in lignes if l["Groupe"] == groupe
                                 and l["Racine suggérée"] == "Oui"), 1)
        distinctes = fe.lire_xlsx(chemin, "Entités distinctes")
        self.assertTrue(all(l["Groupe"].startswith("E-") for l in distinctes))

    def test_comparaison_ancien_classeur(self):
        dossier = tempfile.mkdtemp()
        ancien = os.path.join(dossier, "ancien.xlsx")
        lignes_anciennes = [l for v in ("G-001", "G-084") for l in [
            {"Groupe": v, "Niveau": "Certain", "ID HubSpot": f.id, "Nom de l’entreprise": f.nom,
             "Racine suggérée": "Oui" if i == 0 else ""}
            for i, f in enumerate(REFERENCES[v])]]
        de.ecrire_classeur(ancien, [], [("Doublons", ["Groupe", "Niveau", "Racine suggérée",
                                                      "ID HubSpot", "Nom de l’entreprise"],
                                         lignes_anciennes)], datetime.datetime(2026, 1, 1))
        comparaison = de.comparer_ancien(ancien, self.analyse.lignes(),
                                         {f.id for f in self.analyse.fiches})
        par_groupe = {l["Ancien groupe"]: l for l in comparaison}
        self.assertEqual(par_groupe["G-001"]["Évolution"], "inchangé")
        self.assertEqual(par_groupe["G-084"]["Nouveau niveau"], "Certain partiel (2/3 fiches)")
        self.assertIn("Entités distinctes", par_groupe["G-084"]["Détail par fiche"])


def fe_entetes():
    return ["Groupe", "Niveau", "Raison du rapprochement", "Racine suggérée", "ID HubSpot",
            "Lien HubSpot", "Nom de l’entreprise", "Ville", "SIREN", "SIRET", "Téléphone",
            "Nom de domaine", "Type de décideur", "Propriétaire", "Contacts associés",
            "Date de création"]


class LectureHubSpot(unittest.TestCase):
    """Lecture par l'API (réponses simulées) : pagination, propriétaires, propriétés."""

    def test_entreprises_et_proprietaires(self):
        client = de.ClientDetection("jeton")
        appels = []
        reponses = {
            "/crm/v3/owners?limit=500&archived=false": {"results": [
                {"id": "22", "firstName": "Gautier", "lastName": "Blanc"}]},
            "/crm/v3/owners?limit=500&archived=true": {"results": [
                {"id": "11", "firstName": "Domitille", "lastName": "Denier"}]},
        }

        def requete(methode, chemin, corps=None):
            appels.append(chemin)
            if chemin.startswith("/crm/v3/objects/companies?"):
                if "after=" not in chemin:
                    return {"results": [{"id": "1", "properties": {
                        "name": "ABBAYE DE LA COUDRE", "city": "Laval Cedex",
                        "phone": "332 43 02 85 85", "domain": "abbaye-coudre.fr",
                        "hubspot_owner_id": "11", "num_associated_contacts": "2",
                        "num_associated_deals": "1", "createdate": "2021-07-09T12:48:00Z"}}],
                        "paging": {"next": {"after": "1"}}}
                return {"results": [{"id": "2", "properties": {
                    "name": "EURL ABBAYE DE LA COUDRE", "city": "LAVAL", "hubspot_owner_id": "22",
                    "num_associated_deals": None}}]}
            return reponses[chemin]
        client.requete = requete
        proprietaires = client.proprietaires()
        fiches = [de.fiche_depuis_hubspot(b, proprietaires)
                  for b in client.lister("companies", de.PROPRIETES, journal=lambda *a: 0)]
        self.assertEqual([f.id for f in fiches], ["1", "2"])
        self.assertEqual(fiches[0].proprietaire, "Domitille DENIER (Deactivated User)")
        self.assertIs(fiches[0].proprietaire_actif, False)
        self.assertEqual(fiches[1].proprietaire, "Gautier BLANC")
        self.assertEqual((fiches[0].nb_contacts, fiches[0].nb_transactions), (2, 1))
        self.assertEqual(fiches[1].nb_transactions, 0)
        self.assertTrue(all("properties=name,domain,website" in c for c in appels
                            if "companies?" in c))
        self.assertEqual(de.evaluer(*fiches).niveau, CERTAIN)


class Export(unittest.TestCase):

    def test_colonnes_absentes(self):
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "export.csv")
        with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["ID de fiche d'informations", "Nom de l'entreprise",
                        "Nom de domaine de l'entreprise", "Ville"])
            w.writerow(["14939896642", "ABBAYE DE LA COUDRE", "abbaye-coudre.fr", "Laval"])
            w.writerow(["6543335692", "EURL ABBAYE DE LA COUDRE", "4313.co", "LAVAL CEDEX"])
        fiches, absentes = de.lire_export(chemin)
        self.assertEqual([f.id for f in fiches], ["14939896642", "6543335692"])
        self.assertIn("phone", absentes)
        self.assertIn("siren", absentes)
        self.assertNotIn("city", absentes)
        self.assertEqual(de.main(["--depuis-export", chemin, "-o",
                                  os.path.join(dossier, "sortie.xlsx")]), 0)
        lignes = fe.lire_xlsx(os.path.join(dossier, "sortie.xlsx"))
        self.assertEqual({l["Niveau"] for l in lignes}, {CERTAIN})

    def test_sans_colonne_ville(self):
        a = ent("1", "Abbaye X", "", tel="0243000000")
        b = ent("2", "ABBAYE X", "", tel="0243000000")
        p = de.evaluer(a, b, de.Contexte(villes_connues=False))
        self.assertFalse(p.ville_absente)
        self.assertEqual(p.niveau, CERTAIN)


# --------------------------------------------------------------------------- #
# Contacts (section 6)
# --------------------------------------------------------------------------- #

def contact(id, prenom="", nom="", email="", creation="2024-01-01", **kw):
    return dc.FicheContact(id=id, prenom=prenom, nom=nom, email=email, creation=creation, **kw)


class Contacts(unittest.TestCase):

    def test_b_c1_prenoms_incompatibles(self):
        a = contact(1, "Mélissa", "Belkadi", "nadia.gowsy@century21.fr")
        b = contact(2, "Nadia", "Gowsy", "nadia.gowsy@century21.fr")
        p = dc.evaluer(a, b)
        self.assertTrue(p.bloquee)
        self.assertEqual(p.niveau, "")
        analyse = dc.AnalyseContacts([a, b], journal=lambda *x: 0).executer()
        self.assertFalse(analyse.doublons)
        self.assertEqual(len(analyse.relecture["bloquees"]), 1)

    def test_cc1_meme_email(self):
        a = contact(1, "Jean", "Dupont", "j.dupont@abc.fr")
        b = contact(2, "J", "DUPONT", "J.Dupont@abc.fr ")
        self.assertIn("CC1", dc.evaluer(a, b).regles)
        self.assertEqual(dc.evaluer(a, b).niveau, CERTAIN)        # initiale admise avec CC1
        c = contact(3, "Jean", "Dupont", "jean@autre.fr", emails_secondaires="j.dupont@abc.fr")
        self.assertEqual(dc.evaluer(a, c).niveau, CERTAIN)        # email secondaire

    def test_adresse_generique(self):
        a = contact(1, "Jean", "Dupont", "contact@ecole.fr")
        b = contact(2, "Marie", "Martin", "contact@ecole.fr")
        p = dc.evaluer(a, b)
        self.assertEqual((p.niveau, p.categorie), ("", "generique"))
        c = contact(3, "Jean", "DUPONT", "contact@ecole.fr")
        self.assertEqual(dc.evaluer(a, c).niveau, PROBABLE)       # jamais Certain
        analyse = dc.AnalyseContacts([a, b], journal=lambda *x: 0).executer()
        self.assertEqual(len(analyse.relecture["generique"]), 1)

    def test_cc2_meme_boite(self):
        a = contact(1, "Mon", "Renov", "monprojetrenov@nantesmetropole.fr")
        b = contact(2, "", "RENOV", "monprojetrenov@nantesmetrople.fr")
        self.assertIn("CC2", dc.evaluer(a, b).regles)
        b.nom_net = ""
        b._normes = None
        self.assertEqual(dc.evaluer(a, b).niveau, PROBABLE)       # nom absent

    def test_cc3_et_jamais_certain(self):
        a = contact(1, "Sophie", "Duchêne", "sophie.duchene@gmail.com", id_entreprise="55")
        b = contact(2, "sophie", "DUCHENE", "s.duchene@yahoo.fr", id_entreprise="55")
        self.assertEqual(dc.evaluer(a, b).regles, ["CC3"])
        b = contact(2, "Sophie", "Duchene", "s.duchene@yahoo.fr")
        p = dc.evaluer(a, b)
        self.assertEqual(p.niveau, PROBABLE)                      # webmails seuls
        self.assertIn("webmails seulement", dc.raison(p))
        b = contact(2, "Sophie", "Duchene", "s.duchene@archimeet.fr", id_entreprise="77")
        self.assertEqual(dc.evaluer(a, b).niveau, PROBABLE)       # entreprises différentes
        b = contact(2, "S", "Duchene", "s.duchene@yahoo.fr", id_entreprise="55")
        self.assertEqual(dc.evaluer(a, b).niveau, PROBABLE)       # initiale
        b = contact(2, "Sophie", "Duchene", mobile="06 33 56 45 94")
        a.mobile = "+33 6 33 56 45 94"
        a._normes = None
        self.assertEqual(dc.evaluer(a, b).niveau, CERTAIN)        # même mobile

    def test_noms(self):
        a = contact(1, "Pensabene", "Allonzo", "a@x.fr", id_entreprise="1")
        b = contact(2, "Allonzo", "PENSABENE", "b@x.fr", id_entreprise="1")
        self.assertEqual(dc.evaluer(a, b).niveau, CERTAIN)        # prénom et nom inversés
        a = contact(1, "M.", "BORIES", telephone="+33680180121")
        b = contact(2, "Philippe", "Bories", mobile="06 80 18 01 21")
        self.assertEqual(dc.evaluer(a, b).niveau, PROBABLE)       # nom incomplet
        a = contact(1, "Sylvie", "Kalyniak", id_entreprise="9")
        b = contact(2, "", "Sylvie KALYNIAK", id_entreprise="9")
        self.assertEqual(dc.evaluer(a, b).niveau, CERTAIN)        # nom complet dans un champ

    def test_nom_courant(self):
        fiches = [contact(i, "Pierre", "Martin", "p{}@x{}.fr".format(i, i)) for i in range(5)]
        analyse = dc.AnalyseContacts(fiches, journal=lambda *x: 0).executer()
        self.assertFalse(analyse.doublons)
        analyse = dc.AnalyseContacts(fiches[:2], journal=lambda *x: 0).executer()
        self.assertEqual([g.niveau for g in analyse.doublons], [PROBABLE])

    def test_racine_et_compatibilite_avec_la_fusion(self):
        a = contact(1, "Sophie", "Duchene", "sophie.duchene@archimeet.fr", en_echec=True,
                    id_entreprise="55", activite="2026-01-01")
        b = contact(2, "Sophie", "DUCHENE", "s.duchene@archimeet.fr", id_entreprise="55",
                    creation="2023-01-01")
        c = contact(3, "Marc", "Durand", "marc@x.fr", emails_secondaires="m.durand@y.fr")
        d = contact(4, "Marc", "Durand", "m.durand@y.fr", mobile="0612345678")
        analyse = dc.AnalyseContacts([a, b, c, d], journal=lambda *x: 0).executer()
        groupes = {frozenset(analyse.fiches[m].id for m in g.membres): g for g in analyse.doublons}
        g = groupes[frozenset(("1", "2"))]
        self.assertEqual(analyse.fiches[g.racine].id, "2")          # email non en échec
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "doublons_contacts.xlsx")
        lignes = analyse.lignes()
        horodatage = datetime.datetime(2026, 10, 15)
        de.ecrire_classeur(chemin, dc.synthese(analyse, "test", horodatage),
                           [("Doublons", dc.COLONNES, lignes["Doublons"])], horodatage)
        lues = fc.lire_liste(chemin)
        self.assertEqual(len(lues), 4)
        self.assertEqual(lues[0]["Autres emails"] or lues[1]["Autres emails"] or
                         lues[2]["Autres emails"] or lues[3]["Autres emails"], "m.durand@y.fr")
        compteur = fc.traiter(chemin, os.path.join(dossier, "rapport.csv"), journal=lambda *x: 0)
        self.assertEqual(compteur, {"SIMULATION": 2})


# --------------------------------------------------------------------------- #
# Modifications des scripts de fusion (section 9)
# --------------------------------------------------------------------------- #

def entreprise(id, creation="2022", contacts=0, dates=None, **valeurs):
    return fe.Entreprise(id, valeurs, dates=dates, nb_contacts=contacts, creation=creation)


class FusionEntreprises(unittest.TestCase):

    def test_filet_familles(self):
        a = entreprise("1", name="OGEC MARMOUTIER", city="Tours", phone="0247883535")
        b = entreprise("2", name="Mairie - Marmoutier", city="Tours", phone="0247883535")
        plan = fe.planifier("G-187", [a, b], fe.PROPRIETES_DEFAUT)
        self.assertEqual(plan["statut"], "À VÉRIFIER")
        self.assertTrue(plan["remarques"][0].startswith("familles incompatibles : "))

    def test_filet_chainage(self):
        a = entreprise("1", contacts=3, name="ABBAYE NOTRE DAME DES GARDES",
                       city="SAINT GEORGES DES GARDES", phone="+332 41 29 57 10",
                       domain="www.abbayedesgardes.fr")
        b = entreprise("2", name="abbayedesgardes.fr", phone="332 41 29 57 10",
                       domain="abbayedesgardes.fr")
        c = entreprise("3", name="SARL DE L'ABBAYE NOTRE DAME DES GARDES",
                       city="SAINT GEORGES DES GARDES")
        plan = fe.planifier("G-003", [a, b, c], fe.PROPRIETES_DEFAUT)
        self.assertEqual(plan["statut"], "À VÉRIFIER")
        self.assertTrue(plan["remarques"][0].startswith("lien indirect"))
        self.assertEqual(fe.planifier("G-003", [a, b], fe.PROPRIETES_DEFAUT)["statut"],
                         "À FUSIONNER")

    def test_valeurs_contradictoires(self):
        a = entreprise("1", "2021", contacts=2, name="Abbaye X", city="Laval", phone="0243000000",
                       type_de_decideur="Centralisé", categorie_d_actifs_majoritaires="Habitat",
                       dates={"type_de_decideur": "2021", "categorie_d_actifs_majoritaires": "2021"})
        b = entreprise("2", "2024", name="ABBAYE X", city="LAVAL", phone="02 43 00 00 00",
                       type_de_decideur="Décentralisé",
                       categorie_d_actifs_majoritaires="Ne sait pas encore")
        plan = fe.planifier("G-1", [a, b], fe.PROPRIETES_DEFAUT)
        self.assertEqual(plan["cibles"]["type_de_decideur"], "Décentralisé")   # la plus récente
        self.assertIn("Type de décideur contradictoire : Décentralisé / Centralisé "
                      "(retenu : Décentralisé)", plan["remarques"])
        self.assertFalse(any("Catégorie" in r for r in plan["remarques"]))      # inconnue

    def test_nom_sans_forme_juridique(self):
        eurl = entreprise("1", contacts=3, name="EURL ABBAYE DE LA COUDRE", city="Laval")
        abbaye = entreprise("2", name="ABBAYE DE LA COUDRE", city="LAVAL CEDEX")
        plan = fe.planifier("G-001", [eurl, abbaye], fe.PROPRIETES_DEFAUT)
        self.assertIs(plan["racine"], eurl)
        self.assertEqual(plan["cibles"]["name"], "ABBAYE DE LA COUDRE")
        self.assertIn("name", plan["changements"])
        self.assertEqual(fe.nom_sans_forme_juridique("SARL Cabinet Dupont", [eurl]),
                         "SARL Cabinet Dupont")

    def test_meme_site(self):
        self.assertFalse(fe.meme_site("vendee.gouv.fr", "loire-atlantique.gouv.fr"))
        self.assertTrue(fe.meme_site("https://etu.univ-amu.fr/x", "univ-amu.fr"))
        self.assertTrue(fe.meme_site("intranet.apei.fr", "apei.fr"))

    def test_gt3(self):
        self.assertEqual(fe.ids_gt3_suspects({"123", "8881957130", "2022255"}),
                         ["123", "2022255"])
        f = fe.Entreprise("10", {}, ids_transactions=["555"])
        self.assertTrue(fe.dans_gt3(f, {"10"}))
        ancien, fe.GT3_CONTIENT = fe.GT3_CONTIENT, "transactions"
        try:
            self.assertFalse(fe.dans_gt3(f, {"10"}))
            self.assertTrue(fe.dans_gt3(f, {"555"}))
        finally:
            fe.GT3_CONTIENT = ancien


class FusionContacts(unittest.TestCase):

    def test_lecture_xlsx_niveau_certain(self):
        dossier = tempfile.mkdtemp()
        chemin = os.path.join(dossier, "doublons_contacts.xlsx")
        lignes = [{"Groupe": "G-001", "Niveau": n, "ID HubSpot": i, "Prénom": "Jean",
                   "Nom": "Dupont", "E-mail": "j.dupont@x.fr"}
                  for n, i in ((CERTAIN, "1"), (CERTAIN, "2"), (PROBABLE, "3"))]
        de.ecrire_classeur(chemin, [], [("Doublons", dc.COLONNES, lignes)],
                           datetime.datetime(2026, 1, 1))
        self.assertEqual([l["ID HubSpot"] for l in fc.lire_liste(chemin)], ["1", "2"])

    def test_filet_lien_certain(self):
        a = fc.Fiche(1, prenom="Sébastien", nom="Fesq", email="s.fesq@5-cinq.com",
                     id_entreprise="1", creation="2025")
        b = fc.Fiche(2, prenom="Sébastien", nom="Fesq", email="s.fesq@avramova.org",
                     id_entreprise="2", creation="2026")
        plan = fc.planifier("G-008", [a, b])
        self.assertEqual(plan["statut"], "À VÉRIFIER")
        self.assertIn("lien indirect", plan["remarques"][0])
        b.id_entreprise = "1"
        b._normes = None
        self.assertEqual(fc.planifier("G-008", [a, b])["statut"], "À FUSIONNER")


class CycleHebdomadaire(unittest.TestCase):
    """Relecture et usage récurrent : colonne Décision, mémoire des décisions, colonne Nouveau."""

    def setUp(self):
        self.dossier = tempfile.mkdtemp()
        self.export = os.path.join(self.dossier, "export.csv")
        self.ecrire_export([
            ["1", "ABBAYE DE LA COUDRE", "abbaye-coudre.fr", "Laval", "02 43 02 85 85"],
            ["2", "EURL ABBAYE DE LA COUDRE", "", "LAVAL CEDEX", "+33 2 43 02 85 85"],
            ["3", "Abbaye de la Coudre", "abbaye-coudre.fr", "Laval", ""],
            ["4", "Lycée Saint Louis", "", "Brest", "02 98 00 00 00"],
        ])

    def ecrire_export(self, lignes):
        with open(self.export, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["ID de fiche d'informations", "Nom de l'entreprise",
                        "Nom de domaine de l'entreprise", "Ville", "Numéro de téléphone"])
            w.writerows(lignes)

    def detecter(self, nom, *options):
        sortie = os.path.join(self.dossier, nom)
        self.assertEqual(de.main(["--depuis-export", self.export, "-o", sortie] + list(options)),
                         0)
        return sortie

    def relire(self, chemin, ecartees, sortie):
        """Comme un relecteur : « Ne pas fusionner » sur les lignes des fiches écartées."""
        import openpyxl
        wb = openpyxl.load_workbook(chemin)
        ws = wb["Doublons"]
        entetes = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[entetes.index("ID HubSpot")].value in ecartees:
                row[entetes.index("Décision")].value = "Ne pas fusionner"
        wb.save(sortie)
        return sortie

    def test_colonne_decision(self):
        chemin = self.detecter("semaine1.xlsx")
        lignes = fe.lire_xlsx(chemin)
        self.assertEqual({l["Décision"] for l in lignes if l["Niveau"] == CERTAIN}, {"Fusionner"})
        import openpyxl
        ws = openpyxl.load_workbook(chemin)["Doublons"]
        self.assertEqual(len(ws.data_validations.dataValidation), 1)

    def test_fiche_ecartee_a_la_relecture(self):
        semaine1 = self.detecter("semaine1.xlsx")
        groupe = [l for l in fe.lire_xlsx(semaine1) if l["Niveau"] == CERTAIN]
        self.assertEqual({l["ID HubSpot"] for l in groupe}, {"1", "2", "3"})
        os.makedirs(os.path.join(self.dossier, "relus"))
        relu = self.relire(semaine1, {"2"}, os.path.join(self.dossier, "relus", "s1.xlsx"))
        # Fusion : la fiche 2 est retirée, les fiches 1 et 3 sont fusionnées.
        rapport = os.path.join(self.dossier, "rapport.csv")
        self.assertEqual(fe.traiter(relu, rapport, journal=lambda *a: 0), {"SIMULATION": 1})
        with open(rapport, encoding="utf-8-sig") as f:
            ligne = next(csv.DictReader(f, delimiter=";"))
        self.assertEqual(ligne["Fiches du groupe"], "1 ; 3")
        self.assertIn("écartées à la relecture : 2", ligne["Remarques"])
        # Semaine suivante : la décision est mémorisée, 2 n'est plus proposée avec 1 et 3.
        self.assertEqual(de.lire_decisions([os.path.join(self.dossier, "relus")],
                                           journal=lambda *a: 0),
                         {frozenset(("2", "1")), frozenset(("2", "3"))})
        semaine2 = self.detecter("semaine2.xlsx", "--ancien", relu, "--decisions",
                                 os.path.join(self.dossier, "relus"))
        lignes = fe.lire_xlsx(semaine2)
        self.assertEqual({l["ID HubSpot"] for l in lignes}, {"1", "3"})
        self.assertEqual({l["Nouveau"] for l in lignes}, {""})
        self.assertFalse(fe.lire_xlsx(semaine2, "Entités distinctes"))
        synthese = {l[0]: l[1] for l in iter_synthese(semaine2)}
        self.assertEqual(synthese["Paires écartées à une relecture précédente"], 2)
        self.assertEqual(synthese["Groupes nouveaux depuis l’analyse précédente"], 0)

    def test_memoire_des_decisions(self):
        semaine1 = self.detecter("semaine1.xlsx")
        relu = self.relire(semaine1, {"2"}, os.path.join(self.dossier, "relu.xlsx"))
        memoire = os.path.join(self.dossier, "memoire", "decisions_semaine1.csv")
        os.makedirs(os.path.dirname(memoire))
        sortie = io.StringIO()
        with contextlib.redirect_stdout(sortie):
            self.assertEqual(fe.main([relu, "--hors-ligne", "--memoriser", memoire,
                                      "--rapport", os.path.join(self.dossier, "r.csv")]), 0)
        self.assertIn("2 paires ajoutées", sortie.getvalue())
        self.assertEqual(fc.memoriser_decisions(relu, memoire), 0)       # déjà mémorisées
        silence = lambda *a: None
        attendu = {frozenset(("1", "2")), frozenset(("2", "3"))}
        self.assertEqual(de.lire_decisions([os.path.dirname(memoire)], journal=silence), attendu)
        self.assertEqual(de.lire_decisions([memoire], objet="contact", journal=silence), set())
        semaine2 = self.detecter("semaine2.xlsx", "--decisions", os.path.dirname(memoire))
        self.assertEqual({l["ID HubSpot"] for l in fe.lire_xlsx(semaine2)}, {"1", "3"})

    def test_groupe_entierement_ecarte(self):
        semaine1 = self.detecter("semaine1.xlsx")
        relu = self.relire(semaine1, {"1", "2", "3"}, os.path.join(self.dossier, "relu.xlsx"))
        self.assertEqual(fe.traiter(relu, os.path.join(self.dossier, "r.csv"),
                                    journal=lambda *a: 0), {fc.ECARTE: 1})
        self.assertEqual(len(de.lire_decisions([relu], journal=lambda *a: 0)), 3)

    def test_groupes_nouveaux(self):
        semaine1 = self.detecter("semaine1.xlsx")
        self.ecrire_export([
            ["1", "ABBAYE DE LA COUDRE", "abbaye-coudre.fr", "Laval", "02 43 02 85 85"],
            ["2", "EURL ABBAYE DE LA COUDRE", "", "LAVAL CEDEX", "+33 2 43 02 85 85"],
            ["3", "Abbaye de la Coudre", "abbaye-coudre.fr", "Laval", ""],
            ["4", "Lycée Saint Louis", "", "Brest", "02 98 00 00 00"],
            ["5", "LYCEE SAINT-LOUIS", "", "BREST", "0298000000"],
        ])
        semaine2 = self.detecter("semaine2.xlsx", "--ancien", semaine1)
        nouveaux = {l["ID HubSpot"] for l in fe.lire_xlsx(semaine2) if l["Nouveau"] == "Oui"}
        self.assertEqual(nouveaux, {"4", "5"})

    def test_contacts(self):
        a = contact(1, "Jean", "Dupont", "j.dupont@abc.fr")
        b = contact(2, "Jean", "DUPONT", "j.dupont@abc.fr")
        analyse = dc.AnalyseContacts([a, b], journal=lambda *x: 0,
                                     distinctes={frozenset(("1", "2"))}).executer()
        self.assertFalse(analyse.doublons)
        self.assertFalse(analyse.relecture["bloquees"])
        self.assertEqual(analyse.ecartees_relecture, 1)


class OutilsRoutine(unittest.TestCase):

    def test_reponses_des_connecteurs(self):
        import base64
        import json
        import outils_routine as outils
        dossier = tempfile.mkdtemp()
        classeur = os.path.join(dossier, "c.xlsx")
        de.ecrire_classeur(classeur, [], [("Doublons", ["Groupe"], [{"Groupe": "G-001"}])],
                           datetime.datetime(2026, 1, 1))
        with open(classeur, "rb") as f:
            contenu = base64.b64encode(f.read()).decode()
        reponse = os.path.join(dossier, "r.json")
        with open(reponse, "w") as f:
            json.dump([{"type": "text", "text": json.dumps({"content": contenu, "title": "c"})}], f)
        sortie = os.path.join(dossier, "copie.xlsx")
        self.assertEqual(outils.fichier(reponse, sortie), "c")
        self.assertEqual(fe.lire_xlsx(sortie), [{"Groupe": "G-001"}])
        with open(reponse, "w") as f:                           # copie abîmée : refusée
            json.dump({"content": contenu[:-40] + "A" * 40}, f)
        with self.assertRaises((Exception, SystemExit)):
            outils.fichier(reponse, sortie)
        with open(reponse, "w") as f:
            json.dump({"values": [["ID A", "ID B", "Objet"], ["1", "2", "entreprise"]]}, f)
        memoire = os.path.join(dossier, "memoire.csv")
        self.assertEqual(outils.valeurs(reponse, memoire), 1)
        self.assertEqual(de.lire_decisions([memoire], journal=lambda *a: 0),
                         {frozenset(("1", "2"))})
        self.assertEqual(outils.lignes(memoire), [["1", "2", "entreprise"]])


def iter_synthese(chemin):
    import openpyxl
    return openpyxl.load_workbook(chemin)["Synthèse"].iter_rows(values_only=True)


if __name__ == "__main__":
    unittest.main()
