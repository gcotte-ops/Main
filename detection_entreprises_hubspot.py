#!/usr/bin/env python3
"""
Détection des entreprises HubSpot en doublon : produit doublons_entreprises.xlsx, dont
l'onglet « Doublons » est lu par fusion_entreprises_hubspot.py (lignes de niveau « Certain »).

La détection est en LECTURE SEULE : elle n'écrit jamais dans HubSpot.

Ordre imposé de la chaîne de dédoublonnage :
  1. Détection des entreprises (ce script), puis relecture du classeur.
  2. Fusion des entreprises : fusion_entreprises_hubspot.py doublons_entreprises.xlsx
  3. NOUVELLE extraction des contacts, puis détection des contacts
     (detection_contacts_hubspot.py) : une fois les ID d'entreprise stabilisés par la fusion,
     « même entreprise associée » devient un lien fiable.
  4. Fusion des contacts : fusion_doublons_hubspot.py doublons_contacts.xlsx

Liens entre deux fiches (valeurs normalisées, voir normalisation.py) : même nom, nom hors
parenthèses, nom équivalent (sans forme juridique, type d'établissement ni mots vides), nom à
une faute près, nom inclus, fiche sans nom ou nommée par son domaine ; même domaine, même hôte,
domaine proche ; même téléphone ; même ville ; même SIREN ; « Centralisé » sur les deux fiches.
Liens forts : même domaine, même téléphone, même ville.

  Bloquée    SIREN différents (B1) ou familles incompatibles (B2, voir familles.py) : la paire
             n'est dans aucun groupe, elle est listée dans « Entités distinctes ».
  Certain    C1 nom identique ou équivalent + un lien fort, sans contradiction ;
             C2 nom générique : deux liens forts ; C3 fiche sans nom ou nommée par son domaine
             + même hôte + même téléphone ou même ville ; C4 même SIREN + « Centralisé » sur
             les deux fiches ; C5 nom à une faute près + même SIREN ; C6 domaines différents
             tolérés avec deux liens forts.
             Jamais Certain : nom à une faute près sans SIREN identique, nom inclus, domaine
             seulement proche, villes différentes, ville absente sans autre lien fort,
             domaines différents avec un seul lien fort.
  Probable   nom identique, équivalent, à une faute près ou inclus + un lien fort ou une ville
             absente ; fiche sans nom ou nommée par son domaine + domaine identique ou proche,
             ou même téléphone ; noms différents + même domaine + même téléphone.

Groupes (règle stricte, aucune fusion par chaînage) : dans un groupe Certain, toutes les
paires sont Certain. Sinon, autour de la racine suggérée, on garde les fiches Certain avec
toutes les fiches déjà retenues ; les autres vont dans un groupe Probable « lien indirect
(chaînage) » avec la racine. Les groupes Probable se forment de la même manière sur les
fiches restantes.

Onglets de relecture, jamais fusionnés : « Filiales et établissements », « Homonymes à
vérifier », « SIREN commun », « Entités distinctes ».

Exemples :
  python detection_entreprises_hubspot.py                                  # lecture HubSpot
  python detection_entreprises_hubspot.py --depuis-export entreprises.csv  # sans l'API
  python detection_entreprises_hubspot.py --depuis-mcp base_entreprises.jsonl --nouvelles-depuis 7
                       # base lue par le connecteur HubSpot ; seulement les fiches des 7 derniers jours
  python detection_entreprises_hubspot.py --ancien doublons_potentiels_entreprises_hubspot.xlsx
  python detection_entreprises_hubspot.py --ancien relus/precedent.xlsx --decisions relus/

Relecture : colonne « Décision » de l'onglet Doublons (« Fusionner » sur les groupes Certain).
« Ne pas fusionner » sur une ligne retire la fiche de son groupe ; avec --decisions, ces choix
sont mémorisés et les fiches ne sont plus jamais proposées ensemble. Avec --ancien, la colonne
« Nouveau » signale les groupes absents du classeur précédent (relecture hebdomadaire).

Accès HubSpot : variable d'environnement HUBSPOT_TOKEN = jeton d'une application privée avec
les droits crm.objects.companies.read et crm.objects.owners.read (lecture seule).
Écriture du classeur : openpyxl (pip install openpyxl), utilisé uniquement pour cela.
"""

import argparse
import csv
import datetime
import io
import itertools
import os
import re
import sys
import urllib.parse
import zipfile
from xml.etree import ElementTree

import familles
import fusion_doublons_hubspot as fc
import fusion_entreprises_hubspot as fe
import normalisation as nz
from fusion_doublons_hubspot import (PORTAIL_DEFAUT, ClientHubSpot, ErreurHubSpot, date_iso,
                                     normaliser, verifier_jeton)

FICHIER_SORTIE = "doublons_entreprises.xlsx"
LIEN_FICHE = "https://app-eu1.hubspot.com/contacts/{portail}/record/0-2/{id}"

PROPRIETES = ["name", "domain", "website", "phone", "city", "zip", "address", "siren", "siret",
              "type_de_decideur", "hubspot_owner_id", "num_associated_contacts",
              "num_associated_deals", "createdate", "industry"]

# Nom générique : nom distinctif de 6 caractères ou moins, ou nom équivalent porté par des
# fiches d'au moins 3 villes différentes dans toute la base.
LONGUEUR_NOM_GENERIQUE = 6
VILLES_NOM_GENERIQUE = 3

# Une clé de blocage partagée par plus de fiches que cela (standard téléphonique commun,
# domaine d'un réseau...) ne sert pas à générer de paires : la paire doit partager une autre clé.
TAILLE_MAX_BLOC = 300
# Fiches d'un même ensemble lié au-delà desquelles on ne compare plus toutes les paires.
TAILLE_MAX_COMPOSANTE = 80

CERTAIN, PROBABLE = "Certain", "Probable"
FORTS = ("meme_domaine", "meme_telephone", "meme_ville")
NOMS_LIES = ("meme_nom", "nom_parentheses", "nom_equivalent")

COLONNES = ["Groupe", "Niveau", "Raison du rapprochement", "Racine suggérée", "ID HubSpot",
            "Lien HubSpot", "Nom de l’entreprise", "Ville", "SIREN", "SIRET", "Téléphone",
            "Nom de domaine", "Type de décideur", "Propriétaire", "Contacts associés",
            "Date de création", "Transactions associées", "Dans la GT 3.0",
            "Liens avec la racine", "Nouveau", "Décision"]

# Relecture : choix proposés dans la colonne Décision de l'onglet Doublons.
CHOIX_DECISION = ["Fusionner", "Ne pas fusionner"]
RELUE = "Relecture – fiches déclarées distinctes (« Ne pas fusionner »)"

# (catégorie, préfixe, onglet, niveau affiché)
ONGLETS_RELECTURE = [("filiale", "F-", "Filiales et établissements", "À vérifier"),
                     ("homonyme", "H-", "Homonymes à vérifier", "À vérifier"),
                     ("siren", "S-", "SIREN commun", "Hors liste"),
                     ("bloquees", "E-", "Entités distinctes", "Bloqué")]

# Colonnes d'un export HubSpot (ou de l'ancien classeur), comparées sans accents ni casse.
ALIAS_EXPORT = {
    "id": ["id de fiche d'informations", "record id", "id hubspot", "id", "company id",
           "id de l'entreprise", "hubspot id"],
    "name": ["nom de l'entreprise", "nom", "company name", "name"],
    "domain": ["nom de domaine de l'entreprise", "nom de domaine", "company domain name",
               "domain"],
    "website": ["url du site web", "site web", "website url", "website"],
    "phone": ["numero de telephone", "telephone", "phone number", "phone"],
    "city": ["ville", "city"],
    "zip": ["code postal", "postal code", "zip"],
    "address": ["adresse postale", "adresse", "street address", "address"],
    "siren": ["siren"],
    "siret": ["siret"],
    "type_de_decideur": ["type de decideur"],
    "hubspot_owner_id": ["proprietaire de l'entreprise", "proprietaire", "company owner",
                         "owner"],
    "num_associated_contacts": ["nombre de contacts associes", "contacts associes",
                                "number of associated contacts", "associated contacts"],
    "num_associated_deals": ["nombre de transactions associees", "transactions associees",
                             "number of associated deals", "associated deals"],
    "createdate": ["date de creation", "create date"],
    "industry": ["secteur d'activite", "secteur", "industry"],
}

# Règles désactivées quand une colonne manque dans l'export.
REGLES_PAR_COLONNE = {
    "city": "même ville, villes différentes, ville absente, fréquence des noms génériques",
    "phone": "même téléphone",
    "siren": "SIREN (avec SIRET) : B1, C4, C5, onglet « SIREN commun »",
    "type_de_decideur": "C4 (même SIREN + « Centralisé »)",
    "num_associated_deals": "racine : fiche avec transactions",
    "num_associated_contacts": "racine : nombre de contacts associés",
    "hubspot_owner_id": "propriétaire",
    "createdate": "racine : fiche la plus ancienne",
}


# --------------------------------------------------------------------------- #
# Fiches
# --------------------------------------------------------------------------- #

def entier(valeur):
    try:
        return int(float(str(valeur if valeur is not None else "").strip() or 0))
    except ValueError:
        return 0


def cle_id(id_fiche):
    """Tri numérique des ID HubSpot."""
    return (0, int(id_fiche), "") if str(id_fiche).isdigit() else (1, 0, str(id_fiche))


class FicheEntreprise:
    """Une fiche entreprise : valeurs brutes (affichage) et formes normalisées (comparaison)."""

    def __init__(self, id, nom="", domaine="", site="", telephone="", ville="", code_postal="",
                 adresse="", siren="", siret="", type_de_decideur="", proprietaire="",
                 proprietaire_actif=None, nb_contacts=0, nb_transactions=None, creation="",
                 secteur="", ids_transactions=()):
        self.id = str(id).strip()
        self.nom = nz.nettoyer(nom)
        self.domaine = nz.nettoyer(domaine)
        self.site = nz.nettoyer(site)
        self.telephone = nz.nettoyer(telephone)
        self.ville = nz.nettoyer(ville)
        self.code_postal = nz.nettoyer(code_postal)
        self.adresse = nz.nettoyer(adresse)
        self.siren_brut = nz.nettoyer(siren)
        self.siret_brut = nz.nettoyer(siret)
        self.type_de_decideur = nz.nettoyer(type_de_decideur)
        self.proprietaire = str(proprietaire or "").strip()
        self.proprietaire_actif = proprietaire_actif
        self.nb_contacts = entier(nb_contacts)
        self.nb_transactions = None if nb_transactions is None else entier(nb_transactions)
        self.creation = date_iso(nz.nettoyer(creation))
        self.secteur = nz.nettoyer(secteur)
        self.ids_transactions = {str(i) for i in ids_transactions}

        self.nom_norm = nz.nom_norm(self.nom)
        self.nom_sans_par = nz.nom_sans_parentheses(self.nom)
        self.nom_equiv = nz.nom_equivalent(self.nom)
        self.ville_norm = nz.ville(self.ville)
        self.nom_dist = nz.nom_distinctif(self.nom_equiv, self.ville_norm)
        self.sans_nom = not self.nom_norm
        self.est_domaine = nz.nom_est_domaine(self.nom)
        self.tel = nz.telephone(self.telephone)
        self.hote = nz.hote_utile(self.domaine) or nz.hote_utile(self.site)
        if not self.hote and self.est_domaine:
            self.hote = nz.hote_utile(self.nom)        # « marmoutier.org » sans domaine saisi
        self.domaine_p = nz.domaine_principal(self.hote) if self.hote else ""
        self.siren = nz.siren_depuis(self.siren_brut, self.siret_brut)
        self.centralise = normaliser(self.type_de_decideur) == "centralise"
        self.etiquettes = familles.etiquettes(self.nom)
        # Nom générique : la longueur ici, la fréquence dans la base (Analyse).
        self.generique = len(nz.compact(self.nom_dist)) <= LONGUEUR_NOM_GENERIQUE
        self.lieux = set()

    def __repr__(self):
        return "FicheEntreprise({} {!r})".format(self.id, self.nom)

    @property
    def valeurs(self):
        """Compatibilité avec fusion_entreprises_hubspot.siren_de()."""
        return {"siren": self.siren_brut, "siret": self.siret_brut}


def fiche_depuis_entreprise(entreprise):
    """Fiche de détection construite à partir d'une fiche lue par la fusion."""
    v = entreprise.valeurs
    return FicheEntreprise(
        id=entreprise.id, nom=v.get("name", ""), domaine=v.get("domain", ""),
        site=v.get("website", ""), telephone=v.get("phone", ""), ville=v.get("city", ""),
        code_postal=v.get("zip", ""), adresse=v.get("address", ""), siren=v.get("siren", ""),
        siret=v.get("siret", ""), type_de_decideur=v.get("type_de_decideur", ""),
        proprietaire=v.get("hubspot_owner_id", ""),
        proprietaire_actif=entreprise.proprietaire_actif, nb_contacts=entreprise.nb_contacts,
        nb_transactions=entreprise.nb_transactions, creation=entreprise.creation,
        secteur=v.get("industry", ""))


# --------------------------------------------------------------------------- #
# Comparaison de deux fiches
# --------------------------------------------------------------------------- #

class Contexte:
    """Informations sur toute la base, utiles à la comparaison."""

    def __init__(self, villes_connues=True, lieux=None):
        self.villes_connues = villes_connues      # faux si l'export n'a pas de colonne Ville
        self.lieux = lieux or set()               # villes de la base (forme des noms)


CONTEXTE_DEFAUT = Contexte()


def lieux_du_nom(nom_equiv, contexte):
    """Villes de la base citées dans un nom équivalent (« essca angers » -> {« angers »})."""
    mots = nom_equiv.split()
    trouves = set()
    for n in range(1, min(5, len(mots)) + 1):
        for i in range(len(mots) - n + 1):
            suite = " ".join(mots[i:i + n])
            if suite in contexte.lieux:
                trouves.add(suite)
    return trouves


def mention_de_lieu(longue, courte, contexte):
    """Nom inclus avec mention de lieu : les mots en plus de la fiche au nom long citent une
    ville connue de la base, différente de la ville (connue) de l'autre fiche."""
    autre_ville = nz.forme_ville(courte.ville_norm)
    if not autre_ville:
        return False
    reste = nz.retirer_suite(longue.nom_equiv, courte.nom_equiv)
    for lieu in lieux_du_nom(reste, contexte):
        # « lille » dans « saint andre lille » : même lieu.
        if not (nz.villes_identiques(lieu, autre_ville) or nz.contient_suite(lieu, autre_ville)):
            return True
    return False


class Paire:
    """Liens, contradictions et niveau d'une paire de fiches."""

    def __init__(self, a, b):
        self.a, self.b = a, b
        self.liens = set()
        self.forts = 0
        self.ville_absente = False
        self.generique = False
        self.blocage = []            # motifs B1 / B2
        self.familles = ""           # motif B2
        self.regles = []             # règles Certain satisfaites (C1...C6)
        self.exclusions = []         # raisons de ne pas être Certain (5.4)
        self.niveau_brut = ""        # niveau sans tenir compte du blocage
        self.niveau = ""             # Certain, Probable ou ""
        self.categorie = ""          # onglet de relecture : filiale, homonyme, siren
        self.lieux_differents = False

    @property
    def bloquee(self):
        return bool(self.blocage)

    def autre(self, fiche):
        return self.b if fiche is self.a else self.a


def evaluer(a, b, contexte=CONTEXTE_DEFAUT):
    """Compare deux fiches (sections 5.1 à 5.5 et 5.8)."""
    p = Paire(a, b)
    L = p.liens

    # Noms (espaces ignorés : « APM 22 » = « APM22 »).
    if a.nom_norm and nz.compact(a.nom_norm) == nz.compact(b.nom_norm):
        L.add("meme_nom")
    elif a.nom_sans_par and nz.compact(a.nom_sans_par) == nz.compact(b.nom_sans_par):
        L.add("nom_parentheses")
    elif a.nom_equiv and nz.compact(a.nom_equiv) == nz.compact(b.nom_equiv):
        L.add("nom_equivalent")
    elif a.nom_equiv and b.nom_equiv:
        if (min(len(a.nom_equiv), len(b.nom_equiv)) >= 8
                and nz.distance_un(a.nom_equiv, b.nom_equiv)):
            L.add("nom_faute")
        if nz.contient_suite(a.nom_equiv, b.nom_equiv):
            L.add("nom_inclus")
    if a.sans_nom or b.sans_nom or a.est_domaine or b.est_domaine:
        L.add("nom_domaine")

    # Coordonnées.
    L.update(nz.comparer_domaines(a.hote, b.hote))
    if a.tel and a.tel == b.tel:
        L.add("meme_telephone")
    if contexte.villes_connues:
        if a.ville_norm and b.ville_norm:
            L.add("meme_ville" if nz.villes_identiques(a.ville_norm, b.ville_norm)
                  else "villes_differentes")
        else:
            p.ville_absente = True
    if a.siren and b.siren:
        L.add("meme_siren" if a.siren == b.siren else "siren_differents")
    if a.centralise and b.centralise:
        L.add("centralise_deux")
    p.forts = sum(1 for f in FORTS if f in L)
    p.generique = a.generique or b.generique
    if contexte.lieux:
        la = a.lieux or lieux_du_nom(a.nom_equiv, contexte)
        lb = b.lieux or lieux_du_nom(b.nom_equiv, contexte)
        p.lieux_differents = bool(la and lb and not la & lb)

    # Règles bloquantes.
    if "siren_differents" in L:
        p.blocage.append("B1 – SIREN différents ({} / {})".format(a.siren, b.siren))
    p.familles = familles.comparer_etiquettes(a.etiquettes, b.etiquettes)
    if p.familles:
        p.blocage.append("B2 – familles incompatibles ({})".format(p.familles))

    # Certain (5.3) et exclusions (5.4).
    nom_lie = any(n in L for n in NOMS_LIES)
    vd, dd = "villes_differentes" in L, "domaines_differents" in L
    if nom_lie:
        minimum = 2 if p.generique else 1
        if p.forts >= minimum and not vd and (not dd or p.forts >= 2):
            p.regles.append("C2" if p.generique else "C6" if dd else "C1")
    if ("nom_domaine" in L and "meme_hote" in L and not vd
            and ("meme_telephone" in L or "meme_ville" in L)):
        p.regles.append("C3")
    if "meme_siren" in L and "centralise_deux" in L:
        p.regles.append("C4")
    if "nom_faute" in L and "meme_siren" in L:
        p.regles.append("C5")
    if "nom_faute" in L and "meme_siren" not in L:
        p.exclusions.append("nom à une faute près sans SIREN identique")
    if "nom_inclus" in L:
        p.exclusions.append("nom seulement inclus dans l’autre")
    if vd:
        p.exclusions.append("villes différentes")
    if p.ville_absente and p.forts == 0:
        p.exclusions.append("ville absente sur une fiche, sans autre lien fort")
    if dd and p.forts == 1:
        p.exclusions.append("domaines différents avec un seul lien fort")

    # Filiale ou établissement : nom inclus avec mention d'un autre lieu.
    filiale = False
    if "nom_inclus" in L and contexte.lieux:
        longue, courte = (a, b) if len(a.nom_equiv) > len(b.nom_equiv) else (b, a)
        filiale = mention_de_lieu(longue, courte, contexte)

    # Probable (5.5).
    nom_proche = nom_lie or "nom_faute" in L or "nom_inclus" in L
    if p.regles and not p.exclusions:
        p.niveau_brut = CERTAIN
    elif filiale:
        p.niveau_brut = ""
    elif (p.regles
          or (nom_proche and (p.forts >= 1 or p.ville_absente))
          or ("nom_domaine" in L and L & {"meme_domaine", "domaine_proche", "meme_telephone"})
          or (not nom_proche and "nom_domaine" not in L
              and "meme_domaine" in L and "meme_telephone" in L)):
        p.niveau_brut = PROBABLE

    # Onglets de relecture (paires sans niveau).
    if not p.niveau_brut:
        meme_distinctif = bool(a.nom_dist) and nz.compact(a.nom_dist) == nz.compact(b.nom_dist)
        coordonnees = vd or dd or (a.tel and b.tel and a.tel != b.tel)
        if "meme_siren" in L and "centralise_deux" not in L:
            p.categorie = "siren"
        elif filiale or (("meme_domaine" in L or "meme_telephone" in L) and vd):
            p.categorie = "filiale"
        elif (nom_lie or "nom_faute" in L or meme_distinctif) and coordonnees:
            p.categorie = "homonyme"
    p.niveau = "" if p.bloquee else p.niveau_brut
    return p


LIBELLES_NOM = [("meme_nom", "Même nom"),
                ("nom_parentheses", "Même nom (mention entre parenthèses différente)"),
                ("nom_equivalent", "Nom équivalent (hors type d’établissement)"),
                ("nom_faute", "Nom à une faute près"),
                ("nom_inclus", "Nom inclus dans l’autre")]


def raison(p):
    """« Nom équivalent (hors type d’établissement) + même téléphone + même ville »."""
    L = p.liens
    base = next((libelle for cle, libelle in LIBELLES_NOM if cle in L), "")
    if base and p.generique and any(n in L for n in NOMS_LIES):
        base += " (nom générique)"
    if not base:
        base = ("Fiche sans nom ou nommée par son domaine" if "nom_domaine" in L
                else "Noms différents")
    liens = []
    if "meme_domaine" in L:
        liens.append("même domaine")
    elif "domaine_proche" in L:
        liens.append("domaine proche")
    for cle, libelle in (("meme_telephone", "même téléphone"), ("meme_ville", "même ville")):
        if cle in L:
            liens.append(libelle)
    if "meme_siren" in L:
        liens.append("même SIREN" + (" (Centralisé sur les deux fiches)"
                                     if "centralise_deux" in L else ""))
    texte = base + "".join(" + " + l for l in liens)
    if base == "Noms différents" and "meme_domaine" in L and "meme_telephone" in L:
        texte += " (changement de nom ?)"
    if p.ville_absente and base not in ("Noms différents",):
        texte += ", ville absente sur une fiche"
    contradictions = [libelle for cle, libelle in (("villes_differentes", "villes différentes"),
                                                   ("domaines_differents", "domaines différents"))
                      if cle in L]
    if contradictions:
        texte += " — " + ", ".join(contradictions)
    return texte


def liens_texte(p):
    """Liens et contradictions d'une fiche avec la racine : « même téléphone, même ville »."""
    L = p.liens
    items = []
    for cle, libelle in (("meme_nom", "même nom"), ("nom_parentheses", "même nom hors parenthèses"),
                         ("nom_equivalent", "nom équivalent"), ("nom_faute", "nom à une faute près"),
                         ("nom_inclus", "nom inclus"),
                         ("nom_domaine", "fiche sans nom ou nommée par son domaine")):
        if cle in L:
            items.append(libelle)
    if "meme_hote" in L:
        items.append("même site")
    elif "meme_domaine" in L:
        items.append("même domaine")
    if "domaine_proche" in L and "meme_hote" not in L:
        items.append("domaine proche")
    for cle, libelle in (("meme_telephone", "même téléphone"), ("meme_ville", "même ville"),
                         ("meme_siren", "même SIREN"),
                         ("centralise_deux", "« Centralisé » sur les deux fiches"),
                         ("siren_differents", "SIREN différents"),
                         ("villes_differentes", "villes différentes"),
                         ("domaines_differents", "domaines différents")):
        if cle in L:
            items.append(libelle)
    if p.ville_absente:
        items.append("ville absente")
    if p.familles:
        items.append("familles incompatibles ({})".format(p.familles))
    if p.lieux_differents:
        items.append("lieux différents dans les noms")
    return ", ".join(items) or "aucun lien direct"


# --------------------------------------------------------------------------- #
# Constitution des groupes (règle stricte, commune aux entreprises et aux contacts)
# --------------------------------------------------------------------------- #

class Groupe:

    def __init__(self, niveau, membres, racine, origine=None):
        self.niveau = niveau
        self.membres = membres         # indices des fiches, racine en premier
        self.racine = racine
        self.origine = origine         # groupe auquel les fiches sont rattachées (chaînage)
        self.issus = []                # groupes dont des fiches de celui-ci ont été écartées
        self.numero = ""


def composantes(noeuds, voisins):
    """Composantes connexes (déterministes) du graphe restreint aux noeuds."""
    noeuds = set(noeuds)
    vus, resultat = set(), []
    for depart in sorted(noeuds):
        if depart in vus:
            continue
        vus.add(depart)
        pile, comp = [depart], []
        while pile:
            x = pile.pop()
            comp.append(x)
            for y in voisins.get(x, ()):
                if y in noeuds and y not in vus:
                    vus.add(y)
                    pile.append(y)
        resultat.append(sorted(comp))
    return resultat


class Constructeur:
    """Groupes Certain puis Probable, sans chaînage (section 5.6).

    paire(i, j)        -> objet avec .niveau, .bloquee et .forts (évalué à la demande)
    racine(indices)    -> indice de la racine suggérée

    1. Chaque composante du graphe des paires Certain donne un groupe Certain : la racine et
       les fiches Certain avec toutes les fiches déjà retenues (le plus de liens forts d'abord).
    2. Les fiches restantes et toutes les fiches non placées forment les groupes Probable, de la
       même manière sur le graphe des paires Certain ou Probable ; ce qui reste d'une composante
       est traité à nouveau.
    3. Une fiche écartée d'un groupe Certain et restée seule va dans un groupe Probable « lien
       indirect (chaînage) » avec la racine de ce groupe (jamais avec une fiche bloquée) ; de
       même pour un groupe Probable si elle a un lien direct avec sa racine.
    """

    def __init__(self, n, paire, racine):
        self.n = n
        self.paire = paire
        self.racine = racine

    def _ordre(self, racine):
        return lambda m: (-self.paire(m, racine).forts, m)

    def _decouper(self, comp, niveaux):
        """(noyau, restantes) : racine + fiches liées à la racine et à toutes les retenues."""
        r = self.racine(comp)
        lie = lambda x, y: self.paire(x, y).niveau in niveaux
        autres = sorted((m for m in comp if m != r), key=self._ordre(r))
        if len(comp) <= TAILLE_MAX_COMPOSANTE and all(
                lie(x, y) for x, y in itertools.combinations(comp, 2)):
            return [r] + autres, []
        noyau = [r]
        for m in autres:
            if lie(m, r) and all(lie(m, o) for o in noyau[1:]):
                noyau.append(m)
        return noyau, [m for m in autres if m not in noyau]

    def _rattacher(self, restantes, racine):
        """Fiches restantes rattachées à la racine (aucune paire bloquée dans le groupe)."""
        retenues = []
        for m in sorted(restantes, key=self._ordre(racine)):
            if self.paire(m, racine).bloquee or any(self.paire(m, o).bloquee for o in retenues):
                continue
            retenues.append(m)
        return retenues

    def construire(self, voisins_certain, voisins_lies):
        groupes, places, origine = [], set(), {}
        # 1. Groupes Certain.
        for comp in composantes(range(self.n), voisins_certain):
            if len(comp) < 2:
                continue
            noyau, restantes = self._decouper(comp, (CERTAIN,))
            g = Groupe(CERTAIN, noyau, noyau[0])
            groupes.append(g)
            places.update(noyau)
            for m in restantes:
                origine.setdefault(m, g)
        # 2. Groupes Probable, sur les fiches non placées.
        file, seules = [], []
        for comp in composantes([i for i in range(self.n) if i not in places], voisins_lies):
            (file if len(comp) >= 2 else seules).extend([comp] if len(comp) >= 2 else comp)
        while file:
            comp = file.pop(0)
            if len(comp) < 2:
                seules.extend(comp)
                continue
            noyau, restantes = self._decouper(comp, (CERTAIN, PROBABLE))
            if len(noyau) < 2:
                seules.extend(noyau)
                file.extend(composantes(restantes, voisins_lies))
                continue
            g = Groupe(PROBABLE, noyau, noyau[0])
            g.issus = list({id(origine[m]): origine[m] for m in noyau if m in origine}.values())
            groupes.append(g)
            places.update(noyau)
            for m in restantes:
                origine.setdefault(m, g)
            file.extend(composantes(restantes, voisins_lies))
        # 3. Fiches restées seules : rattachées à la racine du groupe Certain dont elles ont été
        #    écartées, ou à celle d'un groupe Probable si elles ont un lien direct avec elle.
        par_groupe = {}
        for m in sorted(seules):
            g = origine.get(m)
            if g is None or m in places:
                continue
            if g.niveau == CERTAIN or self.paire(m, g.racine).niveau:
                par_groupe.setdefault(id(g), (g, []))[1].append(m)
        for g, membres in par_groupe.values():
            chaine = self._rattacher(membres, g.racine)
            if chaine:
                groupes.append(Groupe(PROBABLE, [g.racine] + chaine, g.racine, origine=g))
        return groupes


def numeroter(groupes, cle, prefixe):
    """G-001, G-002... triés par nom de la racine (numérotation stable)."""
    groupes.sort(key=lambda g: (cle(g.racine), 0 if g.origine is None else 1,
                                0 if g.niveau == CERTAIN else 1, g.membres))
    largeur = max(3, len(str(len(groupes))))
    for n, g in enumerate(groupes, 1):
        g.numero = "{}{:0{}d}".format(prefixe, n, largeur)
    return groupes


# --------------------------------------------------------------------------- #
# Analyse de la base
# --------------------------------------------------------------------------- #

def cle_racine(fiche, gt3=None):
    """Racine suggérée (5.7), alignée sur choisir_racine() de la fusion : transactions, GT 3.0,
    contacts associés, vrai nom, complétude, propriétaire actif, ancienneté."""
    avec = (fiche.nb_transactions or 0) > 0
    dans_gt3 = avec and gt3 is not None and fe.dans_gt3(fiche, gt3)
    complete = sum(1 for v in (fiche.siren_brut, fiche.siret_brut, fiche.hote, fiche.tel,
                               fiche.ville, fiche.adresse, fiche.secteur,
                               fiche.type_de_decideur) if v)
    return (avec, dans_gt3, fiche.nb_contacts, nz.vrai_nom(fiche.nom), complete,
            fiche.proprietaire_actif is True, [-ord(c) for c in fiche.creation],
            tuple(-x if isinstance(x, int) else x for x in cle_id(fiche.id)))


def cle_nom(fiche):
    return (not fiche.nom_norm, fiche.nom_norm, cle_id(fiche.id))


class AnalyseDoublons:
    """Moteur commun aux entreprises et aux contacts : paires candidates par blocage,
    groupes Certain et Probable sans chaînage, onglets de relecture, paires bloquées.

    À définir : comparer(a, b) -> paire (.niveau, .niveau_brut, .categorie, .bloquee, .forts),
    cle_racine(fiche), cles_blocage(fiche), cle_tri(fiche) ; CATEGORIES = [(catégorie,
    préfixe)] des onglets de relecture."""

    CATEGORIES = []
    PREFIXE_BLOQUEES = "E-"

    def __init__(self, fiches, journal=print, distinctes=None):
        uniques = {}
        for f in fiches:
            uniques.setdefault(f.id, f)
        self.fiches = sorted(uniques.values(), key=lambda f: cle_id(f.id))
        self.journal = journal
        self.avertissements = []
        self._cache = {}
        # Paires déclarées distinctes à une relecture précédente : jamais regroupées.
        self.distinctes = distinctes or set()
        self.ecartees_relecture = 0

    def paire(self, i, j):
        if i > j:
            i, j = j, i
        p = self._cache.get((i, j))
        if p is None:
            p = self._cache[(i, j)] = self.comparer(self.fiches[i], self.fiches[j])
            if frozenset((self.fiches[i].id, self.fiches[j].id)) in self.distinctes:
                p.relue = True
                p.blocage.insert(0, RELUE)
                p.niveau = ""
        return p

    def racine(self, indices):
        return max(indices, key=lambda i: self.cle_racine(self.fiches[i]))

    def rapprochee(self, p):
        """Une paire bloquée est listée si elle aurait eu un niveau ou une catégorie."""
        return bool(p.niveau_brut or p.categorie)

    def categorie_de(self, p):
        """Onglet de relecture d'une paire sans niveau (jamais pour une paire bloquée)."""
        return "" if p.bloquee else p.categorie

    def _cle_groupe(self, i):
        return self.cle_tri(self.fiches[i])

    def candidates(self):
        """Paires qui partagent au moins une clé de blocage."""
        blocs = {}
        for i, f in enumerate(self.fiches):
            for cle in self.cles_blocage(f):
                blocs.setdefault(cle, []).append(i)
        paires, ignorees = set(), []
        for cle in sorted(blocs):
            membres = blocs[cle]
            if len(membres) > TAILLE_MAX_BLOC:
                ignorees.append("{} : {} ({} fiches)".format(cle[0], cle[1], len(membres)))
                continue
            paires.update(itertools.combinations(membres, 2))
        if ignorees:
            self.avertissements.append(
                "{} clés partagées par plus de {} fiches, ignorées pour former les paires (la "
                "paire doit partager une autre clé) : {}{}".format(
                    len(ignorees), TAILLE_MAX_BLOC, ", ".join(ignorees[:15]),
                    "…" if len(ignorees) > 15 else ""))
        return sorted(paires)

    def executer(self):
        self.journal("Comparaison des fiches…")
        candidates = self.candidates()
        voisins_certain, voisins_lies = {}, {}
        for i, j in candidates:
            p = self.paire(i, j)
            if p.niveau in (CERTAIN, PROBABLE):
                voisins_lies.setdefault(i, []).append(j)
                voisins_lies.setdefault(j, []).append(i)
            if p.niveau == CERTAIN:
                voisins_certain.setdefault(i, []).append(j)
                voisins_certain.setdefault(j, []).append(i)
        self.journal("  {} paires comparées".format(len(candidates)))
        self.ecartees_relecture = sum(1 for i, j in candidates
                                      if getattr(self.paire(i, j), "relue", False))

        constructeur = Constructeur(len(self.fiches), self.paire, self.racine)
        self.doublons = numeroter(constructeur.construire(voisins_certain, voisins_lies),
                                  self._cle_groupe, "G-")

        # Onglets de relecture : composantes des paires non bloquées de chaque catégorie.
        self.relecture = {}
        for categorie, prefixe in self.CATEGORIES:
            voisins = {}
            for i, j in candidates:
                if self.categorie_de(self.paire(i, j)) == categorie:
                    voisins.setdefault(i, []).append(j)
                    voisins.setdefault(j, []).append(i)
            groupes = []
            for comp in composantes(voisins, voisins):
                if len(comp) >= 2:
                    r = self.racine(comp)
                    groupes.append(Groupe(categorie, [r] + [m for m in comp if m != r], r))
            self.relecture[categorie] = numeroter(groupes, self._cle_groupe, prefixe)
        self.relecture["bloquees"] = numeroter(self._bloquees(candidates, voisins_lies),
                                               self._cle_groupe, self.PREFIXE_BLOQUEES)
        return self

    def _bloquees(self, candidates, voisins_lies):
        """Paires bloquées qui auraient été rapprochées : directement, ou par une fiche liée
        aux deux (« Département 44 » / « Services de l'État 44 » via « loire-atlantique.fr »)."""
        bloquees = {}
        for i, j in candidates:
            p = self.paire(i, j)
            if p.bloquee and self.rapprochee(p) and not getattr(p, "relue", False):
                bloquees[(i, j)] = None
        for x in sorted(voisins_lies):
            voisins = sorted(set(voisins_lies[x]))
            if len(voisins) > TAILLE_MAX_COMPOSANTE:
                continue
            for y, z in itertools.combinations(voisins, 2):
                p = self.paire(y, z)
                if p.bloquee and not getattr(p, "relue", False) and (y, z) not in bloquees:
                    bloquees[(y, z)] = x
        groupes = []
        for (i, j), via in sorted(bloquees.items()):
            r = self.racine([i, j])
            g = Groupe("bloquees", [r, j if r == i else i], r)
            g.via = via
            groupes.append(g)
        return groupes

    def filtrer_nouvelles(self, seuil):
        """Ne garde que les groupes qui contiennent une fiche créée depuis seuil (date_iso) :
        les fiches récentes sont comparées à toute la base. Renvoie le nombre de fiches
        récentes."""
        nouvelle = lambda g: any(self.fiches[m].creation >= seuil for m in g.membres)
        gardes = [g for g in self.doublons if nouvelle(g)]
        # Un groupe ancien n'est pas listé (il ne serait pas à fusionner cette fois) ; un groupe
        # gardé qui y renvoie (chaînage) le désigne par le nom de sa racine.
        for g in self.doublons:
            if g not in gardes:
                f = self.fiches[g.racine]
                nom = (" ".join(x for x in (getattr(f, "prenom", ""), f.nom) if x)
                       or getattr(f, "email", "") or f.id)
                g.numero = "« {} » (groupe plus ancien, non listé)".format(nom)
        self.doublons = numeroter(gardes, self._cle_groupe, "G-")
        prefixes = dict(self.CATEGORIES, bloquees=self.PREFIXE_BLOQUEES)
        for categorie, groupes in self.relecture.items():
            self.relecture[categorie] = numeroter([g for g in groupes if nouvelle(g)],
                                                  self._cle_groupe, prefixes[categorie])
        return sum(1 for f in self.fiches if f.creation >= seuil)

    def compteurs(self, onglets):
        """Chiffres de la synthèse ; onglets : [(catégorie, nom de l'onglet)]."""
        c = {}
        for niveau in (CERTAIN, PROBABLE):
            groupes = [g for g in self.doublons if g.niveau == niveau]
            c[niveau] = (len(groupes), len({m for g in groupes for m in g.membres}))
        c["chainage"] = sum(1 for g in self.doublons if g.origine is not None)
        for categorie, onglet in onglets:
            groupes = self.relecture[categorie]
            c[onglet] = (len(groupes), sum(len(g.membres) for g in groupes))
        return c


class Analyse(AnalyseDoublons):
    """Détection des entreprises en doublon sur toute la base."""

    CATEGORIES = [(c, p) for c, p, _, _ in ONGLETS_RELECTURE[:3]]

    def __init__(self, fiches, gt3=None, villes_connues=True, journal=print, distinctes=None):
        super().__init__(fiches, journal, distinctes)
        self.gt3 = gt3
        lieux = {nz.forme_ville(f.ville_norm) for f in self.fiches}
        self.contexte = Contexte(villes_connues, {l for l in lieux if len(nz.compact(l)) >= 4})
        for f in self.fiches:
            f.lieux = lieux_du_nom(f.nom_equiv, self.contexte)
        self._noms_generiques()

    def _noms_generiques(self):
        villes = {}
        for f in self.fiches:
            if f.nom_equiv and f.ville_norm:
                villes.setdefault(nz.compact(f.nom_equiv), set()).add(f.ville_norm)
        for f in self.fiches:
            if len(villes.get(nz.compact(f.nom_equiv), ())) >= VILLES_NOM_GENERIQUE:
                f.generique = True

    def comparer(self, a, b):
        return evaluer(a, b, self.contexte)

    def cle_racine(self, fiche):
        return cle_racine(fiche, self.gt3)

    def cle_tri(self, fiche):
        return cle_nom(fiche)

    def cles_blocage(self, f):
        cles = set()
        if f.nom_equiv:
            mots = f.nom_equiv.split()
            cles.add(("nom", nz.compact(f.nom_equiv)))
            cles.add(("debut", " ".join(mots[:3])))
            if len(mots) >= 2:
                # Fin du nom, pour les noms inclus (« almaviva sante » / « gie almaviva sante »).
                cles.add(("fin", " ".join(mots[-2:])))
        if f.domaine_p:
            cles.add(("domaine", f.domaine_p))
            # Radical du domaine, pour les domaines proches (marmoutier.org / marmoutier.fr).
            radical = f.domaine_p.split(".")[0]
            if len(radical) >= 4:
                cles.add(("radical", radical))
        if f.tel:
            cles.add(("telephone", f.tel))
        if f.siren:
            cles.add(("siren", f.siren))
        return cles

    # Lignes des onglets ------------------------------------------------------

    def lignes(self, portail=PORTAIL_DEFAUT):
        """{onglet: [lignes]} prêts à écrire."""
        f = self.fiches
        resultat = {"Doublons": []}
        for g in self.doublons:
            r = g.racine
            if g.origine is not None:
                motif = "Lien indirect (chaînage) avec {}".format(g.origine.numero)
                directs = [raison(self.paire(m, r)) for m in g.membres[1:]
                           if self.paire(m, r).niveau]
                if directs:
                    motif += " : " + " ; ".join(dict.fromkeys(directs))
            else:
                motif = " ; ".join(dict.fromkeys(raison(self.paire(m, r))
                                                 for m in g.membres[1:]))
                if g.issus:
                    motif += " — lien indirect (chaînage) avec " + ", ".join(
                        sorted(i.numero for i in g.issus))
            for m in g.membres:
                liens = "(racine)" if m == r else liens_texte(self.paire(m, r))
                if g.origine is not None and m != r and not self.paire(m, r).niveau:
                    liens = "lien indirect, via le groupe {} ; ".format(g.origine.numero) + liens
                resultat["Doublons"].append(self.ligne(f[m], g.numero, g.niveau, motif, m == r,
                                                       liens, portail))
        for categorie, _, onglet, niveau in ONGLETS_RELECTURE:
            lignes = resultat[onglet] = []
            for g in self.relecture[categorie]:
                r = g.racine
                if categorie == "bloquees":
                    p = self.paire(*g.membres)
                    motif = " ; ".join(p.blocage)
                    if g.via is not None:
                        motif += " — rapprochement indirect via « {} » ({})".format(
                            f[g.via].nom or f[g.via].domaine or f[g.via].id, f[g.via].id)
                    elif p.niveau_brut or p.categorie:
                        motif += " — rapprochement écarté : " + raison(p)
                else:
                    motif = " ; ".join(dict.fromkeys(
                        raison(self.paire(x, y)) for x, y in itertools.combinations(g.membres, 2)
                        if self.paire(x, y).categorie == categorie))
                for m in g.membres:
                    liens = "(racine)" if m == r else liens_texte(self.paire(m, r))
                    lignes.append(self.ligne(f[m], g.numero, niveau, motif, m == r, liens,
                                             portail))
        return resultat

    def ligne(self, fiche, groupe, niveau, motif, est_racine, liens, portail):
        return {
            "Groupe": groupe, "Niveau": niveau, "Raison du rapprochement": motif,
            "Racine suggérée": "Oui" if est_racine else "", "ID HubSpot": fiche.id,
            "Lien HubSpot": "Ouvrir", "_url": LIEN_FICHE.format(portail=portail, id=fiche.id),
            "Nom de l’entreprise": fiche.nom, "Ville": fiche.ville, "SIREN": fiche.siren_brut,
            "SIRET": fiche.siret_brut, "Téléphone": fiche.telephone,
            "Nom de domaine": fiche.domaine or fiche.site,
            "Type de décideur": fiche.type_de_decideur, "Propriétaire": fiche.proprietaire,
            "Contacts associés": fiche.nb_contacts, "Date de création": fiche.creation[:16],
            "Transactions associées": "" if fiche.nb_transactions is None else fiche.nb_transactions,
            "Dans la GT 3.0": "" if self.gt3 is None else
            "Oui" if fe.dans_gt3(fiche, self.gt3) else "Non",
            "Liens avec la racine": liens, "Nouveau": "",
            "Décision": "Fusionner" if niveau == CERTAIN else "",
        }

    def compteurs(self):
        return super().compteurs([(c, o) for c, _, o, _ in ONGLETS_RELECTURE])


# --------------------------------------------------------------------------- #
# Lecture des données
# --------------------------------------------------------------------------- #

class ClientDetection(ClientHubSpot):
    """Lecture seule : listes d'objets et propriétaires (avec leurs noms)."""

    def lister(self, objet, proprietes, associations=None, journal=print):
        apres, n = None, 0
        while True:
            chemin = "/crm/v3/objects/{}?limit=100&archived=false&properties={}".format(
                objet, ",".join(proprietes))
            if associations:
                chemin += "&associations=" + associations
            if apres:
                chemin += "&after=" + urllib.parse.quote(apres)
            rep = self.requete("GET", chemin)
            for brut in rep.get("results", []):
                yield brut
            n += len(rep.get("results", []))
            if n and n % 5000 < 100:
                journal("  {} fiches lues…".format(n))
            apres = rep.get("paging", {}).get("next", {}).get("after")
            if not apres:
                break

    def proprietaires(self):
        """{id: (« Prénom NOM », actif)} ; {} si le droit owners.read manque."""
        resultat = {}
        try:
            for archive in ("false", "true"):
                apres = None
                while True:
                    chemin = "/crm/v3/owners?limit=500&archived=" + archive
                    if apres:
                        chemin += "&after=" + urllib.parse.quote(apres)
                    rep = self.requete("GET", chemin)
                    for o in rep.get("results", []):
                        nom = " ".join(x for x in ((o.get("firstName") or "").strip(),
                                                   (o.get("lastName") or "").strip().upper()) if x)
                        resultat.setdefault(str(o["id"]), (nom or o.get("email") or str(o["id"]),
                                                           archive == "false"))
                    apres = rep.get("paging", {}).get("next", {}).get("after")
                    if not apres:
                        break
        except ErreurHubSpot as e:
            print("  ! propriétaires non lus ({}) : ID affichés".format(e), file=sys.stderr)
            return {}
        return resultat


def libelle_proprietaire(id_proprietaire, proprietaires):
    """(« Prénom NOM (Deactivated User) », actif)."""
    if not id_proprietaire:
        return "", None
    if id_proprietaire not in proprietaires:
        return id_proprietaire, None
    nom, actif = proprietaires[id_proprietaire]
    return (nom if actif else nom + " (Deactivated User)"), actif


def fiche_depuis_hubspot(brut, proprietaires):
    p = brut.get("properties", {})
    proprietaire, actif = libelle_proprietaire(p.get("hubspot_owner_id") or "", proprietaires)
    transactions = [a["id"] for a in brut.get("associations", {}).get("deals", {})
                    .get("results", [])]
    return FicheEntreprise(
        id=brut["id"], nom=p.get("name"), domaine=p.get("domain"), site=p.get("website"),
        telephone=p.get("phone"), ville=p.get("city"), code_postal=p.get("zip"),
        adresse=p.get("address"), siren=p.get("siren"), siret=p.get("siret"),
        type_de_decideur=p.get("type_de_decideur"), proprietaire=proprietaire,
        proprietaire_actif=actif, nb_contacts=p.get("num_associated_contacts"),
        nb_transactions=p.get("num_associated_deals") or 0,
        creation=p.get("createdate") or brut.get("createdAt"), secteur=p.get("industry"),
        ids_transactions=transactions)


def fiche_depuis_mcp(p):
    """Fiche lue par le connecteur HubSpot (query_crm_data) : propriétés et libellés
    (« hubspot_owner_id_formatted » = nom du propriétaire)."""
    proprietaire = str(p.get("hubspot_owner_id_formatted") or p.get("hubspot_owner_id") or "").strip()
    return FicheEntreprise(
        id=p["hs_object_id"], nom=p.get("name"), domaine=p.get("domain"), site=p.get("website"),
        telephone=p.get("phone"), ville=p.get("city"), code_postal=p.get("zip"),
        adresse=p.get("address"), siren=p.get("siren"), siret=p.get("siret"),
        type_de_decideur=p.get("type_de_decideur"), proprietaire=proprietaire,
        proprietaire_actif=False if "deactivated" in proprietaire.lower() else None,
        nb_contacts=p.get("num_associated_contacts"),
        nb_transactions=p.get("num_associated_deals") or 0,
        creation=p.get("createdate_iso") or "", secteur=p.get("industry"))


def lire_mcp(chemins):
    """Fiches des réponses du connecteur HubSpot enregistrées (voir outils_routine.py)."""
    import outils_routine
    return [fiche_depuis_mcp(p) for p in outils_routine.enregistrements(chemins, "entreprises")]


def seuil_nouvelles(jours, maintenant=None):
    """Date (UTC, format de date_iso) à partir de laquelle une fiche est « nouvelle » : minuit,
    il y a N jours (une routine qui part quelques minutes en retard ne laisse pas de trou)."""
    maintenant = maintenant or datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    return (maintenant - datetime.timedelta(days=jours)).strftime("%Y-%m-%d 00:00:00")


def feuilles_xlsx(chemin):
    """Noms des onglets d'un classeur."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with zipfile.ZipFile(chemin) as z:
        classeur = ElementTree.fromstring(z.read("xl/workbook.xml"))
    return [s.get("name") for s in classeur.iter(ns + "sheet")]


def lire_tableau(chemin):
    """Lignes d'un export CSV, ou de tous les onglets d'un classeur .xlsx (l'ancien classeur
    de doublons, par exemple), sous forme de dictionnaires."""
    if chemin.lower().endswith((".xlsx", ".xlsm")):
        lignes = []
        for feuille in feuilles_xlsx(chemin):
            lignes.extend(fe.lire_xlsx(chemin, feuille))
        return lignes
    with open(chemin, "rb") as f:
        brut = f.read()
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            texte = brut.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    premiere = texte.splitlines()[0] if texte.strip() else ""
    separateur = ";" if premiere.count(";") > premiere.count(",") else ","
    return list(csv.DictReader(io.StringIO(texte), delimiter=separateur))


def colonnes_export(entetes, alias=ALIAS_EXPORT):
    """{propriété: colonne} d'après les alias (sans accents ni casse)."""
    normes = {}
    for e in entetes:
        normes.setdefault(normaliser(str(e)), e)
    trouvees = {}
    for propriete, noms in alias.items():
        for nom in noms:
            if normaliser(nom) in normes:
                trouvees[propriete] = normes[normaliser(nom)]
                break
    return trouvees


def lire_export(chemin, avertir=print):
    """(fiches, colonnes absentes) depuis un export HubSpot ou l'ancien classeur."""
    lignes = lire_tableau(chemin)
    entetes = []
    for l in lignes:
        for e in l:
            if e not in entetes:
                entetes.append(e)
    col = colonnes_export(entetes)
    if "id" not in col or "name" not in col:
        raise SystemExit("Colonnes « ID » et « Nom de l'entreprise » introuvables dans {} "
                         "(colonnes : {})".format(chemin, ", ".join(map(str, entetes))))
    absentes = [p for p in ALIAS_EXPORT if p not in col]
    if "siren" in absentes and "siret" not in absentes:
        absentes.remove("siren")
    if "domain" in absentes and "website" not in absentes:
        absentes.remove("domain")
    fiches = []
    for l in lignes:
        v = {p: l.get(c) for p, c in col.items()}
        id_fiche = str(v.get("id") or "").strip()
        if not re.fullmatch(r"\d+(\.0)?", id_fiche):
            continue
        proprietaire = str(v.get("hubspot_owner_id") or "").strip()
        fiches.append(FicheEntreprise(
            id=id_fiche.split(".")[0], nom=v.get("name"), domaine=v.get("domain"),
            site=v.get("website"), telephone=v.get("phone"), ville=v.get("city"),
            code_postal=v.get("zip"), adresse=v.get("address"), siren=v.get("siren"),
            siret=v.get("siret"), type_de_decideur=v.get("type_de_decideur"),
            proprietaire=proprietaire,
            proprietaire_actif=("deactivated" not in proprietaire.lower()) if proprietaire else None,
            nb_contacts=v.get("num_associated_contacts"),
            nb_transactions=v.get("num_associated_deals") if "num_associated_deals" in col else None,
            creation=str(v.get("createdate") or ""), secteur=v.get("industry")))
    return fiches, absentes


# --------------------------------------------------------------------------- #
# Classeur
# --------------------------------------------------------------------------- #

BLEU_FONCE, BLEU_CLAIR, JAUNE = "07072D", "EAF0F9", "FFC40B"
POLICE_TITRE, POLICE_TEXTE = "Fraunces", "Open Sans"
LARGEURS = {"Groupe": 9, "Niveau": 11, "Raison du rapprochement": 48, "Racine suggérée": 10,
            "ID HubSpot": 15, "Lien HubSpot": 9, "Ville": 18, "SIREN": 12, "SIRET": 17,
            "Téléphone": 17, "Nom de domaine": 26, "Type de décideur": 14, "Propriétaire": 24,
            "Contacts associés": 10, "Date de création": 16, "Transactions associées": 12,
            "Dans la GT 3.0": 9, "Liens avec la racine": 42, "Prénom": 14, "Nom": 18,
            "E-mail": 30, "Entreprise": 28, "Intitulé du poste": 24,
            "Propriétaire du contact": 24, "Autres emails": 30, "Mobile": 17,
            "ID entreprise associée": 15}


def ecrire_classeur(chemin, synthese, onglets, horodatage):
    """synthese : [(type, texte, valeur)] ; onglets : [(titre, colonnes, lignes)].
    Colonnes « ID … » écrites en texte, « Lien HubSpot » en lien « Ouvrir »."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        raise SystemExit("openpyxl est nécessaire pour écrire le classeur : pip install openpyxl")

    texte = Font(name=POLICE_TEXTE, size=11)
    gras = Font(name=POLICE_TEXTE, size=11, bold=True)
    entete_police = Font(name=POLICE_TEXTE, size=11, bold=True, color="FFFFFF")
    entete_fond = PatternFill("solid", fgColor=BLEU_FONCE)
    bande = PatternFill("solid", fgColor=BLEU_CLAIR)
    accent = PatternFill("solid", fgColor=JAUNE)
    lien_police = Font(name=POLICE_TEXTE, size=11, color=BLEU_FONCE, underline="single")

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 16
    for n, (genre, libelle, valeur) in enumerate(synthese, 1):
        a = ws.cell(row=n, column=1, value=libelle)
        if genre == "titre":
            a.font = Font(name=POLICE_TITRE, size=35, bold=True, color=BLEU_FONCE)
            ws.row_dimensions[n].height = 48
        elif genre == "section":
            a.font = Font(name=POLICE_TITRE, size=16, bold=True, color=BLEU_FONCE)
        elif genre == "paragraphe":
            a.font = texte
            a.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=2)
            ws.row_dimensions[n].height = 15 * (1 + len(libelle) // 95)
        else:
            a.font = texte
        if valeur is not None:
            b = ws.cell(row=n, column=2, value=valeur)
            b.font = gras

    for titre, colonnes, lignes in onglets:
        ws = wb.create_sheet(titre)
        for j, nom in enumerate(colonnes, 1):
            c = ws.cell(row=1, column=j, value=nom)
            c.font, c.fill = entete_police, entete_fond
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(j)].width = LARGEURS.get(nom, 30)
        precedent, fond = None, False
        for i, ligne in enumerate(lignes, 2):
            if ligne.get("Groupe") != precedent:
                precedent, fond = ligne.get("Groupe"), not fond
            for j, nom in enumerate(colonnes, 1):
                valeur = ligne.get(nom, "")
                c = ws.cell(row=i, column=j)
                if nom.startswith("ID"):
                    c.value = "" if valeur in (None, "") else str(valeur)
                    c.number_format = "@"
                elif nom == "Lien HubSpot" and ligne.get("_url"):
                    c.value = "Ouvrir"
                    c.hyperlink = ligne["_url"]
                else:
                    c.value = valeur
                c.font = lien_police if nom == "Lien HubSpot" else texte
                if fond:
                    c.fill = bande
                if nom == "Racine suggérée" and valeur == "Oui":
                    c.fill = accent
        ws.freeze_panes = "A2"
        if lignes:
            ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(colonnes)),
                                                  len(lignes) + 1)
        if titre == "Doublons" and "Décision" in colonnes:
            lettre = get_column_letter(colonnes.index("Décision") + 1)
            choix = DataValidation(type="list", allow_blank=True,
                                   formula1='"{}"'.format(",".join(CHOIX_DECISION)))
            choix.add("{0}2:{0}{1}".format(lettre, max(len(lignes), 1) + 1))
            ws.add_data_validation(choix)
    wb.properties.creator = "Alter Watt – détection des doublons HubSpot"
    wb.properties.created = wb.properties.modified = horodatage
    wb.save(chemin)


# --------------------------------------------------------------------------- #
# Comparaison avec l'ancien classeur (test de non-régression)
# --------------------------------------------------------------------------- #

COLONNES_COMPARAISON = ["Ancien groupe", "Ancien niveau", "Nom", "Fiches", "Nouveau niveau",
                        "Nouveaux groupes", "Évolution", "Détail par fiche"]


def placements(lignes_par_onglet):
    """{ID: [(groupe, niveau, onglet)]} d'après les lignes produites."""
    resultat = {}
    for onglet, lignes in lignes_par_onglet.items():
        for l in lignes:
            resultat.setdefault(str(l["ID HubSpot"]), []).append((l["Groupe"], l["Niveau"], onglet))
    return resultat


def comparer_ancien(chemin_ancien, lignes_par_onglet, ids_analyses, feuille="Doublons",
                    colonne_nom="Nom de l’entreprise"):
    """Lignes du tableau ancien / nouveau niveau par groupe de l'ancien classeur."""
    anciennes = fe.lire_xlsx(chemin_ancien, feuille)
    groupes = {}
    for l in anciennes:
        g = str(l.get("Groupe") or "").strip()
        if g and str(l.get("ID HubSpot") or "").strip():
            groupes.setdefault(g, []).append(l)
    places = placements(lignes_par_onglet)
    doublons = {i: [x for x in v if x[2] == "Doublons"] for i, v in places.items()}
    lignes = []
    for g in sorted(groupes):
        rangs = groupes[g]
        ids = [str(l["ID HubSpot"]).strip() for l in rangs]
        ancien = str(rangs[0].get("Niveau") or "").strip()
        nom = next((str(l.get(colonne_nom) or "") for l in rangs
                    if str(l.get("Racine suggérée") or "").lower() == "oui"),
                   str(rangs[0].get(colonne_nom) or ""))
        nouveaux = {}
        for i in ids:
            for groupe, niveau, _ in doublons.get(i, []):
                nouveaux.setdefault((groupe, niveau), set()).add(i)
        complets = [k for k, v in nouveaux.items() if v == set(ids)]
        certains = {i for (gr, nv), v in nouveaux.items() if nv == CERTAIN for i in v}
        probables = {i for (gr, nv), v in nouveaux.items() if nv == PROBABLE for i in v}
        if any(nv == CERTAIN for _, nv in complets):
            nouveau = CERTAIN
        elif certains:
            nouveau = "Certain partiel ({}/{} fiches)".format(len(certains), len(ids))
        elif any(nv == PROBABLE for _, nv in complets):
            nouveau = PROBABLE
        elif probables:
            nouveau = "Probable partiel ({}/{} fiches)".format(len(probables), len(ids))
        else:
            nouveau = "Hors liste"
        if nouveau == ancien:
            evolution = "inchangé"
        else:
            evolution = "{} → {}".format(ancien or "?", nouveau.split(" (")[0])
        details = []
        for i in ids:
            if i not in ids_analyses:
                details.append("{} : absente des données analysées".format(i))
                continue
            ou = ["{} {}".format(gr, nv if on == "Doublons" else "(" + on + ")")
                  for gr, nv, on in places.get(i, [])]
            details.append("{} : {}".format(i, ", ".join(ou) or "aucun groupe"))
        lignes.append({"Ancien groupe": g, "Ancien niveau": ancien, "Nom": nom,
                       "Fiches": len(ids), "Nouveau niveau": nouveau,
                       "Nouveaux groupes": ", ".join(sorted({gr for gr, _ in nouveaux})),
                       "Évolution": evolution, "Détail par fiche": " ; ".join(details)})
    # Nouveaux groupes Certain sans équivalent dans l'ancien classeur.
    anciens_certains = {str(l["ID HubSpot"]).strip() for v in groupes.values() for l in v
                        if str(l.get("Niveau") or "").strip() == CERTAIN}
    nouveaux_certains = {}
    for l in lignes_par_onglet.get("Doublons", []):
        if l["Niveau"] == CERTAIN:
            nouveaux_certains.setdefault(l["Groupe"], []).append(l)
    for g, rangs in sorted(nouveaux_certains.items()):
        if not any(str(l["ID HubSpot"]) in anciens_certains for l in rangs):
            lignes.append({"Ancien groupe": "", "Ancien niveau": "", "Nom": rangs[0].get(colonne_nom, ""),
                           "Fiches": len(rangs), "Nouveau niveau": CERTAIN,
                           "Nouveaux groupes": g, "Évolution": "nouveau groupe Certain",
                           "Détail par fiche": " ; ".join(str(l["ID HubSpot"]) for l in rangs)})
    return lignes


def bilan_comparaison(lignes):
    compte = {}
    for l in lignes:
        compte[l["Évolution"]] = compte.get(l["Évolution"], 0) + 1
    return sorted(compte.items(), key=lambda x: (-x[1], x[0]))


def onglet_doublons(chemin):
    """Onglet des groupes d'un classeur : « Doublons » (ou « Doublons certains », ancien
    classeur des contacts), None s'il n'y en a pas."""
    feuilles = feuilles_xlsx(chemin)
    return next((f for f in ("Doublons", "Doublons certains") if f in feuilles), None)


def paires_groupees(lignes):
    """Paires d'ID réunies dans un même groupe."""
    groupes = {}
    for l in lignes:
        g, i = str(l.get("Groupe") or "").strip(), str(l.get("ID HubSpot") or "").strip()
        if g and i:
            groupes.setdefault(g, []).append(i)
    return {frozenset(p) for ids in groupes.values() for p in itertools.combinations(ids, 2)
            if p[0] != p[1]}


def marquer_nouveaux(lignes_doublons, chemin_ancien):
    """Colonne « Nouveau » = Oui pour les groupes qui réunissent au moins deux fiches jamais
    groupées ensemble dans le classeur précédent. Renvoie le nombre de groupes nouveaux."""
    feuille = onglet_doublons(chemin_ancien)
    anciennes = paires_groupees(fe.lire_xlsx(chemin_ancien, feuille)) if feuille else set()
    par_groupe = {}
    for l in lignes_doublons:
        par_groupe.setdefault(l["Groupe"], []).append(l)
    nouveaux = 0
    for rangs in par_groupe.values():
        ids = [str(l["ID HubSpot"]) for l in rangs]
        if any(frozenset(p) not in anciennes for p in itertools.combinations(ids, 2)):
            nouveaux += 1
            for l in rangs:
                l["Nouveau"] = "Oui"
    return nouveaux


def fichiers_relus(chemins):
    """Classeurs .xlsx (ou CSV) donnés, ou contenus dans les dossiers donnés."""
    fichiers = []
    for chemin in chemins or ():
        if os.path.isdir(chemin):
            fichiers += sorted(os.path.join(chemin, f) for f in os.listdir(chemin)
                               if f.lower().endswith((".xlsx", ".csv")) and not f.startswith("~$")
                               and not f.lower().endswith("_rapport.csv"))
        elif os.path.isfile(chemin):
            fichiers.append(chemin)
    return fichiers


def lire_decisions(chemins, objet="entreprise", journal=print):
    """Paires d'ID déclarées distinctes lors des relectures précédentes, lues dans les
    fichiers de mémoire (--memoriser des scripts de fusion) et dans les classeurs relus
    (onglet Doublons, lignes « Ne pas fusionner »). objet : « entreprise » ou « contact »."""
    paires, lus = set(), 0
    for chemin in fichiers_relus(chemins):
        try:
            if chemin.lower().endswith(".csv"):
                lignes = lire_tableau(chemin)
            else:
                feuille = onglet_doublons(chemin)
                lignes = fe.lire_xlsx(chemin, feuille) if feuille else []
        except (zipfile.BadZipFile, KeyError, SystemExit, UnicodeDecodeError, OSError) as e:
            journal("  ! relecture illisible, ignorée : {} ({})".format(chemin, e))
            continue
        if not lignes:
            continue
        if "ID A" in lignes[0]:                          # mémoire des décisions
            trouvees = {frozenset((str(l["ID A"]).strip(), str(l["ID B"]).strip()))
                        for l in lignes if str(l.get("Objet") or "").strip() == objet}
        elif "Groupe" in lignes[0] and fc.objet_du_classeur(lignes) == objet:
            trouvees = fc.paires_ecartees(lignes)
        else:
            continue
        lus += 1
        paires |= {p for p in trouvees if len(p) == 2}
    journal("Relectures précédentes : {} fichiers, {} paires déclarées distinctes".format(
        lus, len(paires)))
    return paires


# --------------------------------------------------------------------------- #
# Synthèse
# --------------------------------------------------------------------------- #

LECTURE_NIVEAUX = [
    "Certain (seul niveau fusionné par fusion_entreprises_hubspot.py) : même nom, même nom hors "
    "parenthèses ou nom équivalent (sans forme juridique, type d’établissement ni mots vides) + au "
    "moins un lien fort (même domaine, même téléphone, même ville), sans contradiction. Nom "
    "générique (nom distinctif de 6 caractères ou moins, ou porté dans au moins 3 villes) ou "
    "domaines différents : deux liens forts exigés. Aussi : fiche sans nom ou nommée par son "
    "domaine + même site + même téléphone ou même ville ; même SIREN avec « Centralisé » sur les "
    "deux fiches ; nom à une faute près + même SIREN.",
    "Jamais Certain : nom à une faute près sans SIREN identique, nom seulement inclus dans "
    "l’autre, domaine seulement proche, villes différentes, ville absente sans autre lien fort, "
    "domaines différents avec un seul lien fort.",
    "Probable (à vérifier, jamais fusionné automatiquement) : nom identique, équivalent, à une "
    "faute près ou inclus + un lien fort ou une ville absente ; fiche sans nom ou nommée par son "
    "domaine + domaine identique ou proche, ou même téléphone ; noms différents + même domaine + "
    "même téléphone (changement de nom ?).",
    "Règles bloquantes : SIREN différents (B1) ; familles incompatibles (B2) : deux noms portant "
    "chacun un qualificatif de champs différents, en activité (enseignement, collectivité, État, "
    "médico-social, religieux, immobilier…) ou en nature juridique (association ou fondation, "
    "société, public). Exemples : OGEC / Mairie, Foncière / Association. Association et "
    "fondation sont compatibles. Lexique : lexique_familles.csv. Une paire bloquée n’est dans "
    "aucun groupe : elle figure dans « Entités distinctes » avec son motif.",
    "Règle stricte des groupes : une fusion est irréversible, aucun chaînage n’est accepté. Dans "
    "un groupe Certain, toutes les paires de fiches sont Certain. Une fiche liée à une partie "
    "seulement du groupe va dans un groupe Probable « lien indirect (chaînage) », avec la racine "
    "du groupe Certain : à fusionner à la main après vérification.",
    "Règle SIREN : un SIREN commun ne compte comme lien que si les deux fiches ont « Type de "
    "décideur » = Centralisé ; sinon la paire va dans « SIREN commun ».",
    "Filiales et établissements : même marque ou réseau, sites distincts (même domaine ou même "
    "téléphone avec des villes différentes, nom inclus avec mention d’un autre lieu). Homonymes à "
    "vérifier : même nom distinctif avec domaines, villes ou téléphones contradictoires. Ces "
    "onglets ne sont jamais fusionnés.",
    "Neutralisés : domaines factices « 1234.co », punycode « xn-- », plateformes (wixsite, "
    "e-lyco, google…), domaine réduit à un suffixe public (ac-nantes.fr), téléphones « 00 00 00 "
    "00 00 », valeurs au format Excel (=\"…\"), « Ne sait pas encore », « NC », « inconnu ». "
    "Domaine principal : deux derniers labels, trois pour gouv.fr, asso.fr, ac-*.fr… "
    "(loire-atlantique.gouv.fr ≠ loire-atlantique.fr). Noms comparés sans les espaces "
    "(« APM 22 » = « APM22 »).",
    "Racine suggérée : la fiche avec transactions (celle de la GT 3.0 si plusieurs en ont), puis "
    "le plus de contacts associés, un vrai nom, la plus complète, un propriétaire actif, la plus "
    "ancienne — même ordre que fusion_entreprises_hubspot.py.",
    "Liens avec la racine : liens et contradictions de chaque fiche avec la racine ; « lieux "
    "différents dans les noms » signale des noms qui citent des villes différentes (« ESSCA - "
    "Paris » / « ESSCA - Angers »).",
    "Ordre de la chaîne : 1. détection des entreprises, relecture ; 2. fusion des entreprises ; "
    "3. nouvelle extraction puis détection des contacts ; 4. fusion des contacts.",
]

RELECTURE = [
    "Relecture : dans l’onglet Doublons, la colonne Décision vaut « Fusionner » sur les groupes "
    "Certain. Mettre « Ne pas fusionner » sur la ligne d’une fiche la retire de son groupe "
    "(sur toutes les lignes : le groupe entier est écarté). Seuls les groupes Certain sont "
    "fusionnés automatiquement ; un groupe Probable se fusionne à la main dans HubSpot.",
    "Ces décisions sont mémorisées : une fiche marquée « Ne pas fusionner » n’est plus jamais "
    "proposée avec les autres fiches de son groupe (option --decisions des détections "
    "suivantes). La colonne Nouveau signale les groupes absents de l’analyse précédente.",
]


def lignes_releve(analyse, nouveaux=None, recentes=None):
    """Chiffres de la relecture pour la synthèse."""
    lignes = []
    if recentes is not None:
        lignes.append(("texte", "Fiches créées sur la période (seules leurs paires sont listées)",
                       recentes))
    if nouveaux is not None:
        lignes.append(("texte", "Groupes nouveaux depuis l’analyse précédente", nouveaux))
    if analyse.distinctes:
        lignes.append(("texte", "Paires écartées à une relecture précédente",
                       analyse.ecartees_relecture))
    return lignes


def synthese(analyse, source, horodatage, comparaison=None, avertissements=(), releve=()):
    c = analyse.compteurs()
    lignes = [("titre", "Doublons — entreprises HubSpot", None),
              ("texte", "Analyse du {} — source : {}".format(
                  horodatage.strftime("%d/%m/%Y à %H:%M"), source), None),
              ("vide", "", None),
              ("texte", "Entreprises analysées", len(analyse.fiches)),
              ("texte", "Groupes « Certain » (fusionnés)", c[CERTAIN][0]),
              ("texte", "Fiches « Certain »", c[CERTAIN][1]),
              ("texte", "Groupes « Probable »", c[PROBABLE][0]),
              ("texte", "… dont lien indirect (chaînage)", c["chainage"]),
              ("texte", "Fiches « Probable »", c[PROBABLE][1])]
    for _, _, onglet, _ in ONGLETS_RELECTURE:
        lignes.append(("texte", "{} (hors liste) : fiches".format(onglet), c[onglet][1]))
    lignes += list(releve)
    if comparaison:
        lignes += [("vide", "", None), ("section", "Comparaison avec l’ancien classeur", None)]
        lignes += [("texte", evolution, n) for evolution, n in comparaison]
    if avertissements:
        lignes += [("vide", "", None), ("section", "Avertissements", None)]
        lignes += [("paragraphe", a, None) for a in avertissements]
    lignes += [("vide", "", None), ("section", "Relecture", None)]
    lignes += [("paragraphe", t, None) for t in RELECTURE]
    lignes += [("vide", "", None), ("section", "Comment lire les niveaux", None)]
    lignes += [("paragraphe", t, None) for t in LECTURE_NIVEAUX]
    return lignes


# --------------------------------------------------------------------------- #
# Programme principal
# --------------------------------------------------------------------------- #

def charger_gt3(chemin, journal=print):
    """IDs de la GT 3.0 (None si non fournie), avec l'avertissement des ID suspects."""
    if not chemin:
        return None
    gt3 = fe.lire_gt3(chemin)
    suspects = fe.ids_gt3_suspects(gt3)
    if suspects:
        journal("  ! GT 3.0 : ID de moins de 8 chiffres, suspects : " + ", ".join(suspects))
    return gt3


def chemin_gt3_defaut(args_gt3, dossiers):
    if args_gt3:
        return args_gt3
    return next((os.path.join(d, fe.GT3_DEFAUT) for d in dossiers
                 if os.path.isfile(os.path.join(d, fe.GT3_DEFAUT))), None)


def lire_hubspot(jeton, journal=print):
    client = ClientDetection(jeton)
    probleme = verifier_jeton(client, jeton)
    if not probleme:
        try:
            client.requete("GET", "/crm/v3/objects/companies?limit=1")
        except ErreurHubSpot as e:
            probleme = ("Le jeton ne peut pas lire les entreprises ({}) : ajouter "
                        "crm.objects.companies.read à l'application privée.".format(e))
    if probleme:
        raise SystemExit(probleme)
    proprietaires = client.proprietaires()
    associations = "deals" if fe.GT3_CONTIENT == "transactions" else None
    journal("Lecture des entreprises HubSpot…")
    return [fiche_depuis_hubspot(b, proprietaires)
            for b in client.lister("companies", PROPRIETES, associations, journal)]


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Détecte les entreprises HubSpot en doublon (lecture seule) et produit "
                    "le classeur de relecture lu par fusion_entreprises_hubspot.py.")
    parser.add_argument("--depuis-export", metavar="FICHIER",
                        help="export HubSpot (.csv) ou ancien classeur (.xlsx) au lieu de l'API")
    parser.add_argument("--depuis-mcp", nargs="+", metavar="REPONSES",
                        help="réponses du connecteur HubSpot enregistrées (fichiers, dossiers ou "
                             ".jsonl de outils_routine.py collecter) au lieu de l'API")
    parser.add_argument("--nouvelles-depuis", type=int, metavar="JOURS",
                        help="ne lister que les doublons d'une fiche créée ces N derniers jours "
                             "(comparée à toute la base)")
    parser.add_argument("-o", "--sortie", default=FICHIER_SORTIE,
                        help="classeur produit (défaut : {})".format(FICHIER_SORTIE))
    parser.add_argument("--ancien", metavar="CLASSEUR",
                        help="classeur précédent : onglet « Comparaison » et colonne « Nouveau »")
    parser.add_argument("--decisions", nargs="+", metavar="CLASSEUR_OU_DOSSIER",
                        help="classeurs relus (ou dossiers qui les contiennent) : les fiches "
                             "marquées « Ne pas fusionner » ne sont plus proposées ensemble")
    parser.add_argument("--gt3", help="IDs présents dans la GT 3.0 (défaut : {} s'il est dans "
                        "le dossier) ou export CSV de l'onglet GT3".format(fe.GT3_DEFAUT))
    parser.add_argument("--lexique", help="lexique des familles (défaut : lexique_familles.csv)")
    parser.add_argument("--portail", default=PORTAIL_DEFAUT, help="ID du portail HubSpot")
    args = parser.parse_args(argv)

    for chemin in (args.depuis_export, args.ancien, args.gt3, args.lexique):
        if chemin and not os.path.isfile(chemin):
            parser.error("fichier introuvable : « {} »".format(chemin))
    for chemin in args.decisions or ():
        if not os.path.exists(chemin):
            parser.error("classeur ou dossier de relecture introuvable : « {} »".format(chemin))
    if args.lexique:
        familles._LEXIQUE = familles.charger_lexique(args.lexique)
    gt3 = charger_gt3(chemin_gt3_defaut(args.gt3, [".", os.path.dirname(os.path.abspath(
        args.depuis_export or args.sortie))]))

    avertissements = []
    if args.depuis_export and args.depuis_mcp:
        parser.error("--depuis-export et --depuis-mcp sont incompatibles")
    if args.depuis_mcp:
        fiches = lire_mcp(args.depuis_mcp)
        if not fiches:
            parser.error("aucune entreprise dans les réponses du connecteur : {}".format(
                ", ".join(args.depuis_mcp)))
        source = "connecteur HubSpot (MCP)"
        villes_connues = True
    elif args.depuis_export:
        fiches, absentes = lire_export(args.depuis_export)
        source = "export {}".format(os.path.basename(args.depuis_export))
        for colonne in absentes:
            if colonne in REGLES_PAR_COLONNE:
                message = "Colonne « {} » absente de l'export : règles désactivées : {}.".format(
                    colonne, REGLES_PAR_COLONNE[colonne])
                avertissements.append(message)
                print("  ! " + message)
        villes_connues = "city" not in absentes
    else:
        jeton = os.environ.get("HUBSPOT_TOKEN")
        if not jeton:
            parser.error("définir HUBSPOT_TOKEN (jeton d'application privée, lecture seule), "
                         "ou utiliser --depuis-export")
        fiches = lire_hubspot(jeton.strip())
        source = "API HubSpot (portail {})".format(args.portail)
        villes_connues = True
    print("{} entreprises lues".format(len(fiches)))
    print("GT 3.0 : " + ("{} ID".format(len(gt3)) if gt3 is not None else "non fournie"))

    distinctes = lire_decisions(args.decisions) if args.decisions else None
    analyse = Analyse(fiches, gt3=gt3, villes_connues=villes_connues,
                      distinctes=distinctes).executer()
    recentes = None
    if args.nouvelles_depuis is not None:
        seuil = seuil_nouvelles(args.nouvelles_depuis)
        recentes = analyse.filtrer_nouvelles(seuil)
        print("Fiches créées depuis le {} (UTC) : {}".format(seuil[:10], recentes))
    avertissements += analyse.avertissements
    lignes = analyse.lignes(args.portail)
    horodatage = datetime.datetime.now().replace(microsecond=0)

    onglets = [("Doublons", COLONNES, lignes["Doublons"])]
    onglets += [(onglet, COLONNES, lignes[onglet]) for _, _, onglet, _ in ONGLETS_RELECTURE]
    bilan, nouveaux = None, None
    if args.ancien:
        nouveaux = marquer_nouveaux(lignes["Doublons"], args.ancien)
        comparaison = comparer_ancien(args.ancien, lignes, {f.id for f in analyse.fiches},
                                      feuille=onglet_doublons(args.ancien) or "Doublons")
        onglets.append(("Comparaison", COLONNES_COMPARAISON, comparaison))
        bilan = bilan_comparaison(comparaison)
    ecrire_classeur(args.sortie, synthese(analyse, source, horodatage, bilan, avertissements,
                                          lignes_releve(analyse, nouveaux, recentes)),
                    onglets, horodatage)

    c = analyse.compteurs()
    print("\nCertain : {} groupes ({} fiches) ; Probable : {} groupes ({} fiches, dont {} "
          "liens indirects)".format(c[CERTAIN][0], c[CERTAIN][1], c[PROBABLE][0],
                                    c[PROBABLE][1], c["chainage"]))
    for _, _, onglet, _ in ONGLETS_RELECTURE:
        print("{} : {} groupes, {} fiches".format(onglet, *c[onglet]))
    if nouveaux is not None:
        print("Groupes nouveaux depuis l'analyse précédente : {}".format(nouveaux))
    if distinctes:
        print("Paires écartées à une relecture précédente : {}".format(
            analyse.ecartees_relecture))
    if bilan:
        print("Comparaison avec l'ancien classeur : " +
              ", ".join("{} {}".format(n, e) for e, n in bilan))
    print("Classeur : {}".format(args.sortie))
    print("Rien n'a été modifié dans HubSpot. Après relecture : "
          "python fusion_entreprises_hubspot.py {} --hors-ligne".format(args.sortie))
    return 0


if __name__ == "__main__":
    sys.exit(main())
